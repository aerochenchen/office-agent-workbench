import os
import sys
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ============================================================
# 党政机关公文排版工具 (Format Gongwen)
# 依据标准：GB/T 9704-2012《党政机关公文格式》
# 用途：将用户上传的 .docx 文件调整为标准公文格式
# ============================================================

# 正文段落统一版式
BODY_FONT_PT = 16  # 三号
LINE_SPACING_PT = 29  # 固定行距 29 磅
# 首行缩进 4 个半角字符 = 2 个汉字宽度；OOXML firstLineChars 以「字符」百分之一计
FIRST_LINE_INDENT_CHARS = 2


def set_page_margins(doc):
    """设置A4纸页边距与版心"""
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(37)
        section.bottom_margin = Mm(35)
        section.left_margin = Mm(28)
        section.right_margin = Mm(26)


def set_default_font(doc, font_name='仿宋_GB2312', font_name_ascii='Times New Roman', size=Pt(BODY_FONT_PT)):
    """设置文档默认字体为3号仿宋"""
    style = doc.styles['Normal']
    font = style.font
    font.name = font_name
    font.size = size
    r = style.element.rPr
    if r is None:
        r = OxmlElement('w:rPr')
        style.element.append(r)
    rFonts = r.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        r.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), font_name)
    rFonts.set(qn('w:ascii'), font_name_ascii)
    rFonts.set(qn('w:hAnsi'), font_name_ascii)


def set_line_spacing(doc, line_spacing=LINE_SPACING_PT):
    """设置固定行距（默认 29pt）、段前段后 0"""
    style = doc.styles['Normal']
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(line_spacing)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)


def set_first_line_indent_chars(para, chars=FIRST_LINE_INDENT_CHARS):
    """按字符数设置首行缩进（chars 为汉字字符数；4 半角 = 2 汉字）。"""
    p_element = para._element
    pPr = p_element.find(qn('w:pPr'))
    if pPr is None:
        pPr = OxmlElement('w:pPr')
        p_element.insert(0, pPr)
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        pPr.append(ind)
    # 清除可能冲突的绝对缩进，改用字符缩进
    for attr in ('firstLine', 'hanging'):
        key = qn(f'w:{attr}')
        if ind.get(key) is not None:
            del ind.attrib[key]
    ind.set(qn('w:firstLineChars'), str(int(chars * 100)))
    # 同步绝对缩进（三号 16pt × 汉字数），兼容不认 firstLineChars 的阅读器
    ind.set(qn('w:firstLine'), str(int(BODY_FONT_PT * chars * 20)))


def clear_first_line_indent(para):
    """顶格：清除首行缩进。"""
    para.paragraph_format.first_line_indent = Pt(0)
    p_element = para._element
    pPr = p_element.find(qn('w:pPr'))
    if pPr is None:
        return
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        return
    for attr in ('firstLine', 'firstLineChars', 'hanging'):
        key = qn(f'w:{attr}')
        if ind.get(key) is not None:
            del ind.attrib[key]


def apply_body_paragraph_format(para, *, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    """正文统一：两端对齐、段前段后 0、固定行距 29 磅、可选首行缩进。"""
    para.alignment = align
    pf = para.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(LINE_SPACING_PT)
    if indent:
        set_first_line_indent_chars(para, FIRST_LINE_INDENT_CHARS)
    else:
        clear_first_line_indent(para)


def adjust_paragraph_font(para, font_name='仿宋_GB2312', size=Pt(BODY_FONT_PT), bold=False, color=None):
    """调整段落内所有文本的字体、字号、粗细和颜色"""
    for run in para.runs:
        run.font.name = font_name
        run.font.size = size
        run.bold = bold
        if color:
            run.font.color.rgb = color
        r = run._element.rPr
        if r is None:
            r = OxmlElement('w:rPr')
            run._element.append(r)
        rFonts = r.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts')
            r.insert(0, rFonts)
        rFonts.set(qn('w:eastAsia'), font_name)


def insert_red_line(doc, after_paragraph_index, mm_below=4):
    """在指定段落下方插入一条红色分隔线"""
    para = doc.paragraphs[after_paragraph_index]
    p_element = para._element
    pPr = p_element.find(qn('w:pPr'))
    if pPr is None:
        pPr = OxmlElement('w:pPr')
        p_element.insert(0, pPr)
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '12')  # 约0.5mm
    bottom.set(qn('w:space'), str(int(mm_below * 56.7)))
    bottom.set(qn('w:color'), 'FF0000')
    pBdr.append(bottom)
    pPr.append(pBdr)


def format_docx(file_path):
    """对单个docx文件进行全自动公文排版"""
    doc = Document(file_path)

    # 1. 页面设置
    set_page_margins(doc)

    # 2. 默认字体与行距
    set_default_font(doc)
    set_line_spacing(doc)

    # 3. 识别并调整各要素
    body_start = 0
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            apply_body_paragraph_format(para, indent=False)
            continue

        # 发文机关标志（红色小标宋，居中）
        if i == 0 and ('文件' in text or '通知' in text or '部' in text):
            apply_body_paragraph_format(para, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
            adjust_paragraph_font(para, '方正小标宋简体', Pt(22), bold=False, color=RGBColor(0xFF, 0x00, 0x00))
            body_start = i + 1
            continue

        # 标题（2号小标宋，黑色，居中）
        if para.alignment == WD_ALIGN_PARAGRAPH.CENTER and i <= body_start + 3:
            apply_body_paragraph_format(para, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
            adjust_paragraph_font(para, '方正小标宋简体', Pt(22))
            body_start = i + 1
            continue

        # 主送机关（3号仿宋，顶格）
        if '：' in text and i == body_start:
            apply_body_paragraph_format(para, indent=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            adjust_paragraph_font(para, '仿宋_GB2312', Pt(BODY_FONT_PT))
            body_start = i + 1
            continue

        # 一级标题（一、...）：黑体
        if text.startswith('一、') or text.startswith('二、') or text.startswith('三、'):
            apply_body_paragraph_format(para, indent=True)
            adjust_paragraph_font(para, '黑体', Pt(BODY_FONT_PT))
            continue

        # 二级标题（（一）...）：楷体
        if text.startswith('（一）') or text.startswith('（二）'):
            apply_body_paragraph_format(para, indent=True)
            adjust_paragraph_font(para, '楷体_GB2312', Pt(BODY_FONT_PT))
            continue

        # 普通正文：两端对齐 + 首行缩进 4 半角字符
        apply_body_paragraph_format(para, indent=True)
        adjust_paragraph_font(para, '仿宋_GB2312', Pt(BODY_FONT_PT))

    # 4. 在发文字号后插入红色分隔线（简单处理：在最后一个版头元素后插入）
    # 这里采用启发式方法：查找包含"〔"和"〕"的段落，在其后插入线条
    for i, para in enumerate(doc.paragraphs):
        if '〔' in para.text and '〕' in para.text and '号' in para.text:
            insert_red_line(doc, i)
            break

    # 5. 处理附件说明、署名、日期等（保留原格式，后续可细化）
    # 此处仅做基础调整，保证主体内容符合规范

    # 保存文件
    output_path = file_path.replace('.docx', '_formatted.docx')
    doc.save(output_path)
    return output_path


if __name__ == '__main__':
    if len(sys.argv) != 2:
        print('用法: python format_gongwen.py <文件路径>')
        print('例: python format_gongwen.py /path/to/your/document.docx')
        sys.exit(1)

    input_file = sys.argv[1]
    if not os.path.exists(input_file):
        print(f'错误: 文件不存在 - {input_file}')
        sys.exit(1)

    out = format_docx(input_file)
    print(f'排版完成: {out}')
