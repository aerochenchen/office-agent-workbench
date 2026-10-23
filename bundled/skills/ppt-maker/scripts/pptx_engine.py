"""Semantic deck engine: palettes, validation, named layouts, restyle.

Models fill content (type/title/bullets/data). This module owns positions,
type scale, and colors. Closed page types plus optional variants.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

SKILL_DIR = Path(__file__).resolve().parent.parent
PALETTES_DIR = SKILL_DIR / "palettes"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

SLIDE_TYPES = (
    "title",
    "toc",
    "section",
    "bullets",
    "two_column",
    "kpi",
    "table",
    "chart",
    "mixed_chart",
    "process",
    "closing",
    "canvas",
)
TYPE_ALIAS = {"content": "bullets"}
VARIANTS = {
    "bullets": {"cards"},
    "kpi": {"hero"},
    "process": {"chevron"},
    "chart": {"split"},
    "mixed_chart": {"split"},
}
MAX_BULLETS = 6
MAX_EMPHASIS = 2
MAX_KPI = 4
MAX_STEPS = 6
MAX_TABLE_ROWS = 10
TOKEN_KEYS = (
    "primary",
    "secondary",
    "accent",
    "background",
    "ink",
    "muted",
    "on_primary",
)
POSITION_KEYS = ("left", "top", "width", "height")
CHART_TYPES = {
    "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar_h": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
}
SHAPE_TYPES = {
    "rectangle": MSO_SHAPE.RECTANGLE,
    "rounded_rectangle": MSO_SHAPE.ROUNDED_RECTANGLE,
    "oval": MSO_SHAPE.OVAL,
    "diamond": MSO_SHAPE.DIAMOND,
    "right_arrow": MSO_SHAPE.RIGHT_ARROW,
    "chevron": MSO_SHAPE.CHEVRON,
}
BRIEF_DECK_REWRITE = (
    "旧 brief-deck 页纲请改 type：title/section/content→bullets/closing，"
    "并加上 palette（缺省 zhengwu-navy）。不要写 left/top/pt/裸 hex。"
)


def is_semantic_spec(spec: dict) -> bool:
    if not isinstance(spec, dict):
        return False
    if spec.get("palette") is not None:
        return True
    for slide in spec.get("slides") or []:
        if isinstance(slide, dict) and slide.get("type"):
            return True
    return False


def list_palette_ids() -> list[str]:
    if not PALETTES_DIR.is_dir():
        return []
    return sorted(p.stem for p in PALETTES_DIR.glob("*.json"))


def _hex(value: str) -> str:
    text = str(value or "").strip().lstrip("#").upper()
    if len(text) != 6:
        raise ValueError(f"颜色须为 6 位十六进制，收到 {value!r}")
    int(text, 16)
    return text


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(_hex(hex_color))


def _hex_to_rgb(hex_color: str) -> tuple[int, int, int]:
    h = _hex(hex_color)
    return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)


def _rgb_to_hex(r: int, g: int, b: int) -> str:
    return f"{max(0, min(255, int(r))):02X}{max(0, min(255, int(g))):02X}{max(0, min(255, int(b))):02X}"


def _mix(hex_a: str, hex_b: str, t: float) -> str:
    a = _hex_to_rgb(hex_a)
    b = _hex_to_rgb(hex_b)
    return _rgb_to_hex(
        a[0] + (b[0] - a[0]) * t,
        a[1] + (b[1] - a[1]) * t,
        a[2] + (b[2] - a[2]) * t,
    )


def _luma(hex_color: str) -> float:
    r, g, b = _hex_to_rgb(hex_color)
    return (0.2126 * r + 0.7152 * g + 0.0722 * b) / 255.0


def palette_from_primary(primary: str) -> dict[str, str]:
    p = _hex(primary)
    r, g, b = _hex_to_rgb(p)
    is_red = r > 110 and r > g * 1.35 and r > b * 1.35
    secondary = _mix(p, "FFFFFF", 0.38)
    if _luma(secondary) > 0.72:
        secondary = _mix(p, "FFFFFF", 0.22)
    accent = "D4A84B" if is_red else "C4A35A"
    ink = _mix("1A2332", p, 0.12)
    return {
        "id": "custom",
        "name": "单位主色",
        "primary": p,
        "secondary": secondary,
        "accent": accent,
        "background": "F5F7FA",
        "ink": ink,
        "muted": "5C6B7A",
        "on_primary": "FFFFFF",
        "font_title": "微软雅黑",
        "font_body": "微软雅黑",
    }


def load_palette(spec_palette: Any) -> dict[str, str]:
    if spec_palette is None or spec_palette == "":
        spec_palette = "zhengwu-navy"
    if isinstance(spec_palette, dict):
        if spec_palette.get("primary") and not spec_palette.get("secondary"):
            generated = palette_from_primary(spec_palette["primary"])
            generated.update({k: _hex(v) if k in TOKEN_KEYS else v
                              for k, v in spec_palette.items()
                              if v is not None and k != "primary"})
            generated["primary"] = _hex(spec_palette["primary"])
            return generated
        missing = [k for k in ("primary", "secondary", "accent", "background", "ink", "muted")
                   if not spec_palette.get(k)]
        if missing:
            raise ValueError("palette 对象缺少: " + ", ".join(missing))
        out = {k: (_hex(spec_palette[k]) if k in TOKEN_KEYS else spec_palette[k])
               for k in TOKEN_KEYS if spec_palette.get(k)}
        out["id"] = str(spec_palette.get("id") or "custom")
        out["name"] = str(spec_palette.get("name") or "定制")
        out["on_primary"] = _hex(spec_palette.get("on_primary") or "FFFFFF")
        out["font_title"] = str(spec_palette.get("font_title") or "微软雅黑")
        out["font_body"] = str(spec_palette.get("font_body") or "微软雅黑")
        return out
    pid = str(spec_palette).strip()
    path = PALETTES_DIR / f"{pid}.json"
    if not path.is_file():
        known = "、".join(list_palette_ids()) or "(无)"
        raise ValueError(f"未知 palette {pid!r}。可选：{known}")
    data = json.loads(path.read_text(encoding="utf-8"))
    for key in ("primary", "secondary", "accent", "background", "ink", "muted"):
        data[key] = _hex(data[key])
    data["on_primary"] = _hex(data.get("on_primary") or "FFFFFF")
    data.setdefault("font_title", "微软雅黑")
    data.setdefault("font_body", "微软雅黑")
    return data


def palette_token_hexes(palette: dict[str, str]) -> set[str]:
    return {_hex(palette[k]) for k in TOKEN_KEYS if palette.get(k)}


def resolve_color(value: str, palette: dict[str, str]) -> str:
    raw = str(value or "").strip()
    if raw in palette and raw in TOKEN_KEYS:
        return _hex(palette[raw])
    if raw in ("white", "on_primary"):
        return _hex(palette.get("on_primary") or "FFFFFF")
    hexed = _hex(raw)
    allowed = palette_token_hexes(palette)
    if hexed not in allowed:
        raise ValueError(
            f"颜色 {value!r} 不在当前 palette token 中。"
            f"请用 primary/secondary/accent/background/ink/muted/on_primary"
        )
    return hexed


def _slide_type(slide: dict) -> str:
    raw = str(slide.get("type") or "").strip().lower()
    return TYPE_ALIAS.get(raw, raw)


def _items(value: Any) -> list:
    if value is None:
        return []
    if isinstance(value, list):
        return value
    return [value]


def _bullet_text(item: Any) -> tuple[str, int]:
    if isinstance(item, str):
        return item.strip(), 0
    if isinstance(item, dict):
        return str(item.get("text") or "").strip(), int(item.get("level") or 0)
    return str(item).strip(), 0


def _count_bullets(slide: dict) -> int:
    stype = _slide_type(slide)
    if stype == "two_column":
        return len(_items(slide.get("left") or slide.get("bullets"))) + len(
            _items(slide.get("right") or slide.get("right_bullets"))
        )
    if stype == "kpi":
        return len(_items(slide.get("kpis") or slide.get("bullets")))
    if stype == "process":
        return len(_items(slide.get("steps") or slide.get("bullets")))
    if stype in ("toc", "bullets", "closing", "mixed_chart"):
        return len(_items(slide.get("bullets")))
    return 0


def validate_deck(spec: Any) -> dict[str, Any]:
    errors: list[str] = []
    warnings: list[str] = []
    if not isinstance(spec, dict):
        return {"ok": False, "errors": ["deck.json 须为对象"], "warnings": [],
                "slide_count": 0, "palette": None}
    slides = spec.get("slides")
    if not isinstance(slides, list) or not slides:
        errors.append("缺少非空 slides 数组")
        return {"ok": False, "errors": errors, "warnings": warnings,
                "slide_count": 0, "palette": None}

    if not is_semantic_spec(spec):
        errors.append(
            "这是旧版 layout 坐标稿，不能当正式页纲。"
            "请改用 type（title/toc/section/bullets/two_column/kpi/table/"
            "chart/mixed_chart/process/closing）。"
            + BRIEF_DECK_REWRITE
        )
        return {"ok": False, "errors": errors, "warnings": warnings,
                "slide_count": len(slides), "palette": None}

    palette = None
    try:
        palette = load_palette(spec.get("palette"))
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        errors.append(str(exc))

    allowed = "、".join(SLIDE_TYPES)
    for i, slide in enumerate(slides):
        prefix = f"第{i + 1}页"
        if not isinstance(slide, dict):
            errors.append(f"{prefix} 不是对象")
            continue
        stype = _slide_type(slide)
        if stype not in SLIDE_TYPES:
            errors.append(
                f"{prefix} type={slide.get('type')!r} 不在封闭列表：{allowed}。"
                + BRIEF_DECK_REWRITE
            )
            continue
        variant = slide.get("variant")
        if variant:
            allowed_v = VARIANTS.get(stype, set())
            short = str(variant).split(".")[-1]
            if short not in allowed_v:
                errors.append(
                    f"{prefix} variant={variant!r} 不可用。"
                    f"{stype} 可选：{', '.join(sorted(allowed_v)) or '无'}"
                )
        if stype != "canvas":
            for key in POSITION_KEYS:
                if key in slide and isinstance(slide[key], (int, float)):
                    errors.append(f"{prefix} 非 canvas 页禁止写 {key}，由引擎排版")
            for key in ("fill", "color", "background", "size", "font"):
                if key in slide and stype != "canvas":
                    errors.append(f"{prefix} 非 canvas 页禁止写 {key}")
        n_bullets = _count_bullets(slide)
        limit = MAX_KPI if stype == "kpi" else MAX_STEPS if stype == "process" else MAX_BULLETS
        if n_bullets > limit:
            errors.append(f"{prefix} 条目 {n_bullets} 条，超过上限 {limit}")
        emphasis = slide.get("emphasis") or []
        if isinstance(emphasis, list) and len(emphasis) > MAX_EMPHASIS:
            errors.append(f"{prefix} emphasis 最多 {MAX_EMPHASIS} 处")
        if stype == "table":
            rows = slide.get("rows") or []
            if not rows:
                errors.append(f"{prefix} table 页需要 rows")
            elif len(rows) > MAX_TABLE_ROWS:
                errors.append(f"{prefix} 表格超过 {MAX_TABLE_ROWS} 行，请拆页")
        if stype in ("chart", "mixed_chart"):
            chart = slide.get("chart") or {}
            if not isinstance(chart, dict) or not chart.get("categories") or not chart.get("series"):
                errors.append(f"{prefix} 需要 chart.categories 与 chart.series")
            kind = str((chart or {}).get("kind") or (chart or {}).get("type") or "bar")
            if kind not in CHART_TYPES:
                errors.append(f"{prefix} 图表 kind={kind!r}，可选 bar/bar_h/line/pie")
        if stype == "canvas" and palette is not None:
            try:
                _validate_canvas(slide, palette, prefix, errors)
            except ValueError as exc:
                errors.append(f"{prefix} {exc}")

    return {
        "ok": not errors,
        "errors": errors,
        "warnings": warnings,
        "slide_count": len(slides),
        "palette": None if palette is None else palette.get("id"),
    }


def _validate_canvas(slide: dict, palette: dict[str, str], prefix: str, errors: list[str]) -> None:
    for group in ("shapes", "tables", "charts", "images"):
        for item in slide.get(group) or []:
            if not isinstance(item, dict):
                continue
            for key in ("fill", "text_color", "color"):
                if item.get(key):
                    try:
                        resolve_color(str(item[key]), palette)
                    except ValueError as exc:
                        errors.append(f"{prefix} canvas {key}: {exc}")


def _fill_solid(shape, hex_color: str) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(hex_color)
    shape.line.fill.background()


def _set_run(run, text: str, *, size_pt: int, bold: bool = False,
             color: str, font: str, italic: bool = False) -> None:
    run.text = text
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = _rgb(color)
    run.font.name = font


def _textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    box.text_frame.word_wrap = True
    return box


def _blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _paint_bg(slide, hex_color: str) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), SLIDE_W, SLIDE_H)
    _fill_solid(shape, hex_color)


def _left_band(slide, palette: dict[str, str]) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), Inches(0.18), SLIDE_H
    )
    _fill_solid(shape, palette["primary"])


def _footer(slide, label: str, palette: dict[str, str]) -> None:
    box = _textbox(slide, Inches(0.6), Inches(7.05), Inches(12.2), Inches(0.32))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    _set_run(run, (label or "")[:80], size_pt=10, color=palette["muted"],
             font=palette["font_body"])


def _title_bar(slide, title: str, palette: dict[str, str]) -> None:
    head = _textbox(slide, Inches(0.55), Inches(0.38), Inches(12.2), Inches(0.7))
    p = head.text_frame.paragraphs[0]
    run = p.add_run()
    _set_run(run, title or "要点", size_pt=22, bold=True,
             color=palette["primary"], font=palette["font_title"])


def _write_bullets(tf, bullets: list, palette: dict[str, str],
                   emphasis: set[int], *, size_pt: int = 18) -> None:
    items = [_bullet_text(b) for b in bullets]
    items = [(t, lv) for t, lv in items if t][:MAX_BULLETS]
    if not items:
        items = [("（本页要点待补）", 0)]
    for i, (text, level) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = max(0, min(4, level))
        p.space_after = Pt(8)
        run = p.add_run()
        color = palette["accent"] if i in emphasis else palette["ink"]
        _set_run(run, text, size_pt=size_pt, bold=i in emphasis,
                 color=color, font=palette["font_body"])


def _emphasis_set(slide: dict) -> set[int]:
    raw = slide.get("emphasis") or []
    if not isinstance(raw, list):
        return set()
    out = set()
    for item in raw[:MAX_EMPHASIS]:
        try:
            out.add(int(item))
        except (TypeError, ValueError):
            continue
    return out


def _variant(slide: dict, stype: str) -> str:
    raw = str(slide.get("variant") or "").strip()
    if not raw:
        return ""
    if raw.startswith(stype + "."):
        return raw.split(".", 1)[1]
    return raw


def _add_chart(slide, chart: dict, palette: dict[str, str],
               left, top, width, height) -> None:
    data = CategoryChartData()
    data.categories = [str(c) for c in chart.get("categories") or []]
    series = chart.get("series") or {}
    if isinstance(series, dict):
        items = list(series.items())
    else:
        items = [(str(s.get("name") or f"系列{i+1}"), s.get("values") or [])
                 for i, s in enumerate(series)]
    for name, values in items:
        data.add_series(str(name), [float(v) for v in values])
    kind = str(chart.get("kind") or chart.get("type") or "bar")
    frame = slide.shapes.add_chart(
        CHART_TYPES.get(kind, XL_CHART_TYPE.COLUMN_CLUSTERED),
        left, top, width, height, data,
    )
    if chart.get("title"):
        frame.chart.has_title = True
        frame.chart.chart_title.text_frame.text = str(chart["title"])
    colors = [palette["primary"], palette["secondary"], palette["accent"], palette["muted"]]
    for i, ser in enumerate(frame.chart.series):
        color = _rgb(colors[i % len(colors)])
        try:
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = color
        except ValueError:
            pass
        try:
            ser.format.line.color.rgb = color
        except ValueError:
            pass


def _add_table(slide, rows: list, palette: dict[str, str],
               left, top, width, height) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    shape = slide.shapes.add_table(len(rows), cols, left, top, width, height)
    table = shape.table
    for r, row in enumerate(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            text = str(row[c]) if c < len(row) else ""
            cell.text = text
            fill = palette["primary"] if r == 0 else (
                _mix(palette["background"], palette["secondary"], 0.12) if r % 2 else palette["background"]
            )
            ink = palette["on_primary"] if r == 0 else palette["ink"]
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(fill)
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(12 if r else 13)
                    run.font.bold = r == 0
                    run.font.color.rgb = _rgb(ink)
                    run.font.name = palette["font_body"]


def render_title(slide, spec: dict, deck: dict, palette: dict[str, str]) -> None:
    _paint_bg(slide, palette["primary"])
    title = spec.get("title") or deck.get("title") or "汇报"
    box = _textbox(slide, Inches(0.9), Inches(2.35), Inches(11.5), Inches(1.7))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    _set_run(run, str(title), size_pt=36, bold=True,
             color=palette["on_primary"], font=palette["font_title"])
    subtitle = spec.get("subtitle") or deck.get("subtitle") or ""
    if subtitle:
        sub = _textbox(slide, Inches(0.9), Inches(4.2), Inches(11.5), Inches(0.9))
        sp = sub.text_frame.paragraphs[0]
        sr = sp.add_run()
        _set_run(sr, str(subtitle), size_pt=16,
                 color=_mix(palette["on_primary"], palette["secondary"], 0.25),
                 font=palette["font_body"])
    _footer(slide, str(deck.get("title") or ""), palette)


def render_section(slide, spec: dict, deck: dict, palette: dict[str, str]) -> None:
    _paint_bg(slide, palette["background"])
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(0), Inches(2.95), SLIDE_W, Inches(1.55)
    )
    _fill_solid(bar, palette["primary"])
    box = _textbox(slide, Inches(0.9), Inches(3.2), Inches(11.5), Inches(1.1))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    _set_run(run, str(spec.get("title") or "章节"), size_pt=28, bold=True,
             color=palette["on_primary"], font=palette["font_title"])
    _footer(slide, str(deck.get("title") or ""), palette)


def render_bullets(slide, spec: dict, deck: dict, palette: dict[str, str]) -> None:
    _paint_bg(slide, palette["background"])
    _left_band(slide, palette)
    _title_bar(slide, str(spec.get("title") or "要点"), palette)
    bullets = _items(spec.get("bullets"))
    emphasis = _emphasis_set(spec)
    if _variant(spec, "bullets") == "cards":
        items = [_bullet_text(b)[0] for b in bullets if _bullet_text(b)[0]][:MAX_BULLETS]
        if not items:
            items = ["（本页要点待补）"]
        n = len(items)
        gap = 0.18
        card_h = min(0.85, (5.4 - gap * (n - 1)) / max(n, 1))
        top = 1.25
        for i, text in enumerate(items):
            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(0.55), Inches(top + i * (card_h + gap)),
                Inches(12.2), Inches(card_h),
            )
            fill = palette["accent"] if i in emphasis else _mix(
                palette["background"], palette["secondary"], 0.18
            )
            _fill_solid(card, fill)
            tf = card.text_frame
            tf.word_wrap = True
            tf.paragraphs[0].alignment = PP_ALIGN.LEFT
            run = tf.paragraphs[0].add_run()
            ink = palette["on_primary"] if i in emphasis else palette["ink"]
            _set_run(run, text, size_pt=16, bold=i in emphasis,
                     color=ink, font=palette["font_body"])
    else:
        body = _textbox(slide, Inches(0.7), Inches(1.25), Inches(11.9), Inches(5.5))
        _write_bullets(body.text_frame, bullets, palette, emphasis)
    _footer(slide, str(deck.get("title") or ""), palette)


def render_toc(slide, spec: dict, deck: dict, palette: dict[str, str]) -> None:
    _paint_bg(slide, palette["background"])
    _left_band(slide, palette)
    _title_bar(slide, str(spec.get("title") or "汇报提纲"), palette)
    bullets = _items(spec.get("bullets"))
    numbered = []
    for i, item in enumerate(bullets[:MAX_BULLETS], start=1):
        text, _ = _bullet_text(item)
        numbered.append(f"{i:02d}  {text}")
    body = _textbox(slide, Inches(0.7), Inches(1.3), Inches(11.9), Inches(5.4))
    _write_bullets(body.text_frame, numbered, palette, _emphasis_set(spec), size_pt=20)
    _footer(slide, str(deck.get("title") or ""), palette)


def render_two_column(slide, spec: dict, deck: dict, palette: dict[str, str]) -> None:
    _paint_bg(slide, palette["background"])
    _left_band(slide, palette)
    _title_bar(slide, str(spec.get("title") or "对照"), palette)
    left = list(_items(spec.get("left") or spec.get("bullets")))
    right = list(_items(spec.get("right") or spec.get("right_bullets")))
    left_box = _textbox(slide, Inches(0.6), Inches(1.3), Inches(5.8), Inches(5.4))
    right_box = _textbox(slide, Inches(6.8), Inches(1.3), Inches(5.8), Inches(5.4))
    if spec.get("left_title"):
        left.insert(0, {"text": str(spec["left_title"]), "level": 0})
    if spec.get("right_title"):
        right.insert(0, {"text": str(spec["right_title"]), "level": 0})
    _write_bullets(left_box.text_frame, left, palette, set())
    _write_bullets(right_box.text_frame, right, palette, set())
    _footer(slide, str(deck.get("title") or ""), palette)


def render_kpi(slide, spec: dict, deck: dict, palette: dict[str, str]) -> None:
    _paint_bg(slide, palette["background"])
    _left_band(slide, palette)
    _title_bar(slide, str(spec.get("title") or "主要指标"), palette)
    kpis = []
    for item in _items(spec.get("kpis") or spec.get("bullets"))[:MAX_KPI]:
        if isinstance(item, dict):
            kpis.append((str(item.get("label") or item.get("text") or ""),
                         str(item.get("value") or "")))
        elif isinstance(item, (list, tuple)) and len(item) >= 2:
            kpis.append((str(item[0]), str(item[1])))
        else:
            kpis.append((str(item), ""))
    if not kpis:
        kpis = [("待补", "—")]
    hero = _variant(spec, "kpi") == "hero"
    if hero and kpis:
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(1.3),
            Inches(12.2), Inches(2.2),
        )
        _fill_solid(card, palette["primary"])
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        _set_run(run, kpis[0][1] or kpis[0][0], size_pt=36, bold=True,
                 color=palette["on_primary"], font=palette["font_title"])
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        _set_run(r2, kpis[0][0], size_pt=14, color=palette["on_primary"],
                 font=palette["font_body"])
        rest = kpis[1:]
        top = 3.8
    else:
        rest = kpis
        top = 1.4
    if rest:
        n = len(rest)
        width = (12.2 - 0.2 * (n - 1)) / n
        for i, (label, value) in enumerate(rest):
            left = 0.55 + i * (width + 0.2)
            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top),
                Inches(width), Inches(2.4 if not hero else 2.0),
            )
            _fill_solid(card, _mix(palette["background"], palette["secondary"], 0.2))
            tf = card.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            _set_run(run, value or "—", size_pt=24, bold=True,
                     color=palette["primary"], font=palette["font_title"])
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            r2 = p2.add_run()
            _set_run(r2, label, size_pt=13, color=palette["ink"],
                     font=palette["font_body"])
    _footer(slide, str(deck.get("title") or ""), palette)


def render_table(slide, spec: dict, deck: dict, palette: dict[str, str]) -> None:
    _paint_bg(slide, palette["background"])
    _left_band(slide, palette)
    _title_bar(slide, str(spec.get("title") or "一览"), palette)
    _add_table(
        slide, spec.get("rows") or [], palette,
        Inches(0.55), Inches(1.25), Inches(12.2), Inches(5.4),
    )
    _footer(slide, str(deck.get("title") or ""), palette)


def render_chart(slide, spec: dict, deck: dict, palette: dict[str, str]) -> None:
    _paint_bg(slide, palette["background"])
    _left_band(slide, palette)
    _title_bar(slide, str(spec.get("title") or "数据"), palette)
    chart = spec.get("chart") or {}
    split = _variant(spec, "chart") == "split" or spec.get("bullets")
    if split:
        _add_chart(slide, chart, palette, Inches(0.5), Inches(1.25), Inches(7.4), Inches(5.4))
        body = _textbox(slide, Inches(8.1), Inches(1.3), Inches(4.6), Inches(5.3))
        _write_bullets(body.text_frame, _items(spec.get("bullets")), palette,
                       _emphasis_set(spec), size_pt=16)
    else:
        _add_chart(slide, chart, palette, Inches(0.7), Inches(1.25), Inches(11.9), Inches(5.4))
    _footer(slide, str(deck.get("title") or ""), palette)


def render_mixed_chart(slide, spec: dict, deck: dict, palette: dict[str, str]) -> None:
    spec = dict(spec)
    spec["variant"] = spec.get("variant") or "split"
    render_chart(slide, spec, deck, palette)


def render_process(slide, spec: dict, deck: dict, palette: dict[str, str]) -> None:
    _paint_bg(slide, palette["background"])
    _left_band(slide, palette)
    _title_bar(slide, str(spec.get("title") or "步骤"), palette)
    steps = []
    for item in _items(spec.get("steps") or spec.get("bullets"))[:MAX_STEPS]:
        steps.append(_bullet_text(item)[0] or "步骤")
    if not steps:
        steps = ["待补"]
    n = len(steps)
    chevron = _variant(spec, "process") == "chevron"
    width = min(2.1, (12.0 - 0.15 * (n - 1)) / n)
    top = 2.6
    for i, text in enumerate(steps):
        left = 0.55 + i * (width + 0.15)
        shape_kind = MSO_SHAPE.CHEVRON if chevron else MSO_SHAPE.ROUNDED_RECTANGLE
        shp = slide.shapes.add_shape(shape_kind, Inches(left), Inches(top),
                                     Inches(width), Inches(1.5))
        _fill_solid(shp, palette["primary"] if i == 0 else palette["secondary"])
        tf = shp.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        try:
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            shp.text_frame.auto_size = None
        except (AttributeError, TypeError, ValueError):
            pass
        run = tf.paragraphs[0].add_run()
        label = f"{i + 1}  {text}" if not chevron else text
        _set_run(run, label, size_pt=13, bold=True,
                 color=palette["on_primary"], font=palette["font_body"])
    _footer(slide, str(deck.get("title") or ""), palette)


def render_closing(slide, spec: dict, deck: dict, palette: dict[str, str]) -> None:
    copied = dict(spec)
    copied["title"] = spec.get("title") or "汇报完毕"
    render_bullets(slide, copied, deck, palette)


def render_canvas(slide, spec: dict, deck: dict, palette: dict[str, str],
                  confine_path) -> None:
    bg = spec.get("background")
    _paint_bg(slide, resolve_color(bg, palette) if bg else palette["background"])
    if spec.get("title"):
        _title_bar(slide, str(spec["title"]), palette)
    for img in spec.get("images") or []:
        path = confine_path(str(img["path"]))
        kwargs = {}
        if img.get("width"):
            kwargs["width"] = Inches(float(img["width"]))
        if img.get("height"):
            kwargs["height"] = Inches(float(img["height"]))
        slide.shapes.add_picture(
            path, Inches(float(img.get("left", 1))), Inches(float(img.get("top", 1))),
            **kwargs,
        )
    for tbl in spec.get("tables") or []:
        _add_table(
            slide, tbl.get("rows") or [], palette,
            Inches(float(tbl.get("left", 0.6))),
            Inches(float(tbl.get("top", 1.5))),
            Inches(float(tbl.get("width", 12))),
            Inches(float(tbl.get("height", 3))),
        )
    for shp in spec.get("shapes") or []:
        shape = slide.shapes.add_shape(
            SHAPE_TYPES.get(str(shp.get("type") or "rectangle"), MSO_SHAPE.RECTANGLE),
            Inches(float(shp.get("left", 1))),
            Inches(float(shp.get("top", 1))),
            Inches(float(shp.get("width", 2))),
            Inches(float(shp.get("height", 1))),
        )
        if shp.get("fill"):
            _fill_solid(shape, resolve_color(str(shp["fill"]), palette))
        if shp.get("text"):
            tf = shape.text_frame
            tf.word_wrap = True
            run = tf.paragraphs[0].add_run()
            color = resolve_color(str(shp.get("text_color") or "ink"), palette)
            _set_run(run, str(shp["text"]), size_pt=14, color=color,
                     font=palette["font_body"])
    for cht in spec.get("charts") or []:
        chart = dict(cht)
        if "kind" not in chart and cht.get("type") in CHART_TYPES:
            chart["kind"] = cht["type"]
        _add_chart(
            slide, chart, palette,
            Inches(float(cht.get("left", 1))),
            Inches(float(cht.get("top", 2))),
            Inches(float(cht.get("width", 6))),
            Inches(float(cht.get("height", 3.5))),
        )
    _footer(slide, str(deck.get("title") or spec.get("footer") or ""), palette)


RENDERERS = {
    "title": render_title,
    "toc": render_toc,
    "section": render_section,
    "bullets": render_bullets,
    "two_column": render_two_column,
    "kpi": render_kpi,
    "table": render_table,
    "chart": render_chart,
    "mixed_chart": render_mixed_chart,
    "process": render_process,
    "closing": render_closing,
}


def build_semantic(spec: dict, out_path: Path, *, confine_path) -> dict[str, Any]:
    report = validate_deck(spec)
    if not report["ok"]:
        raise ValueError("；".join(report["errors"]))
    palette = load_palette(spec.get("palette"))
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    counts: dict[str, int] = {}
    for raw in spec.get("slides") or []:
        stype = _slide_type(raw)
        slide = _blank(prs)
        if stype == "canvas":
            render_canvas(slide, raw, spec, palette, confine_path)
        else:
            RENDERERS[stype](slide, raw, spec, palette)
        if raw.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(raw["notes"])
        counts[stype] = counts.get(stype, 0) + 1
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return {
        "ok": True,
        "output": str(out_path),
        "slides": sum(counts.values()),
        "counts": counts,
        "palette": palette.get("id"),
    }


def _iter_shapes(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def restyle_presentation(prs: Presentation, palette: dict[str, str]) -> dict[str, Any]:
    """Recolor fills/text/charts/tables to palette. Does not change copy."""
    changed = 0
    slide_w = int(prs.slide_width)
    slide_h = int(prs.slide_height)
    for slide in prs.slides:
        try:
            fill = slide.background.fill
            fill.solid()
            fill.fore_color.rgb = _rgb(palette["background"])
            changed += 1
        except (ValueError, AttributeError, TypeError):
            pass
        for shape in list(_iter_shapes(slide.shapes)):
            try:
                if shape.has_text_frame and not shape.has_table:
                    w, h = int(shape.width), int(shape.height)
                    full = w >= slide_w * 0.92 and h >= slide_h * 0.92
                    band = w <= int(Inches(0.4)) and h >= slide_h * 0.85
                    bar = w >= slide_w * 0.9 and int(Inches(1.0)) <= h <= int(Inches(2.2))
                    if full or band or bar:
                        _fill_solid(shape, palette["primary"])
                        ink = palette["on_primary"]
                    else:
                        ink = None
                    for para in shape.text_frame.paragraphs:
                        for run in para.runs:
                            if not run.text:
                                continue
                            size = run.font.size.pt if run.font.size else 14
                            if ink:
                                color = ink
                            elif size >= 20 and run.font.bold:
                                color = palette["primary"]
                            elif size <= 11:
                                color = palette["muted"]
                            else:
                                color = palette["ink"]
                            run.font.color.rgb = _rgb(color)
                            changed += 1
                if shape.has_table:
                    table = shape.table
                    for r, row in enumerate(table.rows):
                        for cell in row.cells:
                            fill_hex = palette["primary"] if r == 0 else palette["background"]
                            cell.fill.solid()
                            cell.fill.fore_color.rgb = _rgb(fill_hex)
                            ink = palette["on_primary"] if r == 0 else palette["ink"]
                            for para in cell.text_frame.paragraphs:
                                for run in para.runs:
                                    run.font.color.rgb = _rgb(ink)
                            changed += 1
                if shape.has_chart:
                    colors = [
                        palette["primary"],
                        palette["secondary"],
                        palette["accent"],
                        palette["muted"],
                    ]
                    for i, ser in enumerate(shape.chart.series):
                        color = _rgb(colors[i % len(colors)])
                        try:
                            ser.format.fill.solid()
                            ser.format.fill.fore_color.rgb = color
                        except ValueError:
                            pass
                        try:
                            ser.format.line.color.rgb = color
                        except ValueError:
                            pass
                    changed += 1
            except (ValueError, AttributeError, TypeError):
                continue
    return {"ok": True, "palette": palette.get("id"), "touched": changed}
