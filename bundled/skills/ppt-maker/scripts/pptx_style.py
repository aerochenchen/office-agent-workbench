#!/usr/bin/env python3
"""Extract a style pack from a finished .pptx (palette / type scale / chrome).

Does not copy shape coordinates. Output is consumed by the semantic engine
via deck.json `"style": "path.json"` when the user says 以这份 PPT 为
模板 / 参考 / 样板.
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from pathlib import Path

from office_agent.script_policy import confine_to_workspace
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from pptx_engine import _hex, _luma, _mix, load_palette

EMU = 914400.0
BADGE_CHARS = set("0123456789一二三四五六七八九十")


def _in(value) -> float | None:
    if value is None:
        return None
    return round(float(value) / EMU, 4)


def _solid_hex(shape) -> str | None:
    try:
        fill = shape.fill
        if fill.type is None:
            return None
        if "SOLID" in str(fill.type).upper():
            return _hex(str(fill.fore_color.rgb))
    except Exception:
        return None
    return None


def _is_oval(shape) -> bool:
    try:
        name = str(shape.auto_shape_type).upper()
        return "OVAL" in name or "ELLIPSE" in name
    except Exception:
        return False


def _shape_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return (shape.text_frame.text or "").strip()


def _iter_shapes(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _is_gold(hex_color: str) -> bool:
    h = _hex(hex_color)
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    return r >= 170 and g >= 130 and b <= 200 and r + g > b * 2.2


def extract_style(prs: Presentation) -> dict:
    slide_w = _in(prs.slide_width) or 13.33
    slide_h = _in(prs.slide_height) or 7.5
    fills: list[tuple[str, float, int]] = []
    text_colors: Counter[str] = Counter()
    title_fonts: Counter[str] = Counter()
    body_fonts: Counter[str] = Counter()
    cover_sizes: list[float] = []
    page_title_sizes: list[float] = []
    body_sizes: list[float] = []
    caption_sizes: list[float] = []
    has_top_bar = False
    has_page_number = False
    has_badge = False
    has_top_accent = False
    column_hits = 0

    slides = list(prs.slides)
    for idx, slide in enumerate(slides):
        rects: list[tuple[float, float, float, float, str]] = []
        for shape in _iter_shapes(slide.shapes):
            left, top, width, height = (
                _in(shape.left), _in(shape.top),
                _in(shape.width), _in(shape.height),
            )
            hx = _solid_hex(shape)
            if hx and width and height:
                fills.append((hx, width * height, idx))
                if (width >= slide_w * 0.9 and height <= 0.28
                        and (top or 0) <= 0.22 and idx > 0):
                    has_top_bar = True
                if (width >= 1.5 and height >= 2.8
                        and 0.15 <= _luma(hx) <= 0.98):
                    rects.append((left or 0, top or 0, width, height, hx))
                if (height <= 0.14 and width >= 1.4 and (top or 0) >= 1.4
                        and _luma(hx) < 0.7):
                    has_top_accent = True
            text = _shape_text(shape)
            if _is_oval(shape) and text and text[0] in BADGE_CHARS and width and 0.3 <= width <= 1.1:
                has_badge = True
            elif text in BADGE_CHARS and width and 0.35 <= width <= 0.9 and height and 0.35 <= height <= 0.9:
                has_badge = True
            if (text.isdigit() and len(text) <= 3 and width and width <= 1.0
                    and (top or 0) >= slide_h - 0.55 and (left or 0) >= slide_w - 1.6):
                has_page_number = True
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    chunk = (run.text or "").strip()
                    if not chunk:
                        continue
                    size = run.font.size.pt if run.font.size else None
                    font = run.font.name
                    try:
                        if run.font.color and run.font.color.rgb:
                            text_colors[_hex(str(run.font.color.rgb))] += 1
                    except Exception:
                        pass
                    if idx == 0:
                        if size:
                            cover_sizes.append(size)
                        if font:
                            title_fonts[font] += 3 if run.font.bold else 1
                    else:
                        if size and size >= 22:
                            page_title_sizes.append(size)
                            if font:
                                title_fonts[font] += 2 if run.font.bold else 1
                        elif size and size <= 11:
                            caption_sizes.append(size)
                        elif size:
                            body_sizes.append(size)
                            if font:
                                body_fonts[font] += 1
        if len(rects) >= 3:
            widths = sorted(r[2] for r in rects)
            if widths[-1] - widths[0] <= 0.6:
                column_hits += 1

    area_by = Counter()
    for hx, area, _idx in fills:
        area_by[hx] += area
    dark = [(hx, area) for hx, area in area_by.items() if _luma(hx) < 0.42]
    dark.sort(key=lambda item: item[1], reverse=True)
    primary = dark[0][0] if dark else (max(area_by, key=area_by.get) if area_by else "1B3A5C")

    golds = [(hx, area) for hx, area in area_by.items() if _is_gold(hx)]
    golds.sort(key=lambda item: item[1], reverse=True)
    chromatic = [
        (hx, area) for hx, area in area_by.items()
        if hx != primary and 0.18 <= _luma(hx) < 0.85
    ]
    chromatic.sort(key=lambda item: item[1], reverse=True)
    accent = golds[0][0] if golds else (chromatic[0][0] if chromatic else "C4A35A")
    secondary_cands = [hx for hx, _area in chromatic if hx != accent]
    secondary = secondary_cands[0] if secondary_cands else _mix(primary, "FFFFFF", 0.38)

    surfaces = [
        (hx, area) for hx, area in area_by.items()
        if 0.86 <= _luma(hx) <= 0.98 and hx not in {primary, "FFFFFF"}
    ]
    surfaces.sort(key=lambda item: item[1], reverse=True)
    surface = surfaces[0][0] if surfaces else "F0F4F8"

    light_bg = [
        (hx, area) for hx, area in area_by.items() if _luma(hx) >= 0.92
    ]
    background = "FFFFFF"
    if light_bg and light_bg[0][1] > slide_w * slide_h * 0.15:
        background = max(light_bg, key=lambda item: item[1])[0]

    ink_cands = [(hx, n) for hx, n in text_colors.items() if _luma(hx) < 0.4]
    ink = max(ink_cands, key=lambda item: item[1])[0] if ink_cands else "2C3E50"
    muted_cands = [(hx, n) for hx, n in text_colors.items() if 0.4 <= _luma(hx) <= 0.7]
    muted = max(muted_cands, key=lambda item: item[1])[0] if muted_cands else "7F8C8D"
    on_primary = "FFFFFF"

    cover_dark = False
    if slides:
        cover_fills = [(hx, area) for hx, area, idx in fills if idx == 0]
        if cover_fills:
            biggest = max(cover_fills, key=lambda item: item[1])
            cover_dark = _luma(biggest[0]) < 0.45

    def _median(values: list[float], fallback: int) -> int:
        if not values:
            return fallback
        ordered = sorted(values)
        return int(round(ordered[len(ordered) // 2]))

    font_title = title_fonts.most_common(1)[0][0] if title_fonts else "微软雅黑"
    font_body = body_fonts.most_common(1)[0][0] if body_fonts else font_title

    palette = load_palette({
        "id": "extracted",
        "name": "参考稿风格",
        "primary": primary,
        "secondary": secondary,
        "accent": accent,
        "background": background,
        "ink": ink,
        "muted": muted,
        "on_primary": on_primary,
        "font_title": font_title,
        "font_body": font_body,
    })
    return {
        "palette": {k: palette[k] for k in (
            "id", "name", "primary", "secondary", "accent", "background",
            "ink", "muted", "on_primary", "font_title", "font_body",
        )},
        "surface": surface,
        "type_scale": {
            "cover_kicker": _median([s for s in cover_sizes if s < 36], 20),
            "cover_title": max([int(s) for s in cover_sizes] or [36]),
            "cover_meta": _median([s for s in cover_sizes if 14 <= s <= 24], 16),
            "page_title": _median(page_title_sizes, 22),
            "body": _median(body_sizes, 16),
            "caption": _median(caption_sizes, 10),
            "badge": 16,
        },
        "chrome": {
            "cover": "dark-center" if cover_dark else "default",
            "header": "top-bar" if has_top_bar else "left-band",
            "card": "tint-top-accent" if has_top_accent else "plain",
            "footer": "page-number" if has_page_number else "title",
            "number_badge": bool(has_badge),
        },
        "motifs": {
            "toc": "columns" if column_hits else "list",
            "bullets": "cards" if has_top_accent or has_badge else "list",
        },
    }


def _ws(raw: str) -> Path:
    return confine_to_workspace(Path(raw), Path.cwd())


def main(argv=None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser(
        description="从成品 PPT 抽取风格包（配色/字体/骨架），不复制坐标。")
    parser.add_argument("reference", help="参考 .pptx")
    parser.add_argument("output", help="写出 style.json")
    args = parser.parse_args(argv)
    try:
        src = _ws(args.reference)
        out_path = _ws(args.output)
        prs = Presentation(str(src))
        style = extract_style(prs)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_text(json.dumps(style, ensure_ascii=False, indent=2),
                            encoding="utf-8")
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    print(json.dumps({
        "ok": True,
        "output": str(out_path),
        "palette": style["palette"],
        "chrome": style["chrome"],
        "motifs": style["motifs"],
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
