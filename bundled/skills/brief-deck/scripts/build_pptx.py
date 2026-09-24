#!/usr/bin/env python3
"""Build a simple 16:9 formal PPTX from a slides JSON for brief-deck.

Usage:
  python build_pptx.py --slides .office-agent/work/brief-deck/slides.json \\
                       --out 工作成果/汇报演示.pptx
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from office_agent.script_policy import confine_to_workspace

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

# Formal navy-on-paper defaults (aligns with zhengwu-navy spirit; not a palette copy)
NAVY = RGBColor(0x1B, 0x3A, 0x5F)
NAVY_BAND = RGBColor(0x24, 0x4A, 0x73)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x5C, 0x6B, 0x7A)
PAPER = RGBColor(0xF7, 0xF8, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _set_run(run, text: str, *, size_pt: int, bold: bool = False, color: RGBColor = INK) -> None:
    run.text = text
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "微软雅黑"


def _fill_solid(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _add_textbox(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)


def _blank_slide(prs: Presentation):
    layout = prs.slide_layouts[6]  # blank
    return prs.slides.add_slide(layout)


def _paint_paper(slide) -> None:
    shape = slide.shapes.add_shape(
        1,  # MSO_SHAPE.RECTANGLE
        Emu(0),
        Emu(0),
        SLIDE_W,
        SLIDE_H,
    )
    _fill_solid(shape, PAPER)


def _left_band(slide, width=Inches(0.18)) -> None:
    shape = slide.shapes.add_shape(1, Emu(0), Emu(0), width, SLIDE_H)
    _fill_solid(shape, NAVY_BAND)


def _footer(slide, label: str) -> None:
    box = _add_textbox(slide, Inches(0.6), Inches(7.05), Inches(12), Inches(0.35))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    _set_run(run, label[:80], size_pt=10, color=MUTED)


def render_title(slide, title: str, subtitle: str, deck_title: str) -> None:
    band = slide.shapes.add_shape(1, Emu(0), Emu(0), SLIDE_W, Inches(7.5))
    _fill_solid(band, NAVY)
    box = _add_textbox(slide, Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.6))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    _set_run(run, title or deck_title or "汇报", size_pt=36, bold=True, color=WHITE)
    if subtitle:
        sub = _add_textbox(slide, Inches(0.9), Inches(4.2), Inches(11.5), Inches(0.8))
        sp = sub.text_frame.paragraphs[0]
        sr = sp.add_run()
        _set_run(sr, subtitle, size_pt=16, color=RGBColor(0xD6, 0xDE, 0xE8))
    _footer(slide, deck_title)


def render_section(slide, title: str, deck_title: str) -> None:
    _paint_paper(slide)
    bar = slide.shapes.add_shape(1, Emu(0), Inches(3.0), SLIDE_W, Inches(1.5))
    _fill_solid(bar, NAVY)
    box = _add_textbox(slide, Inches(0.9), Inches(3.25), Inches(11.5), Inches(1.0))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    _set_run(run, title or "章节", size_pt=28, bold=True, color=WHITE)
    _footer(slide, deck_title)


def render_content(slide, title: str, bullets: list[str], deck_title: str) -> None:
    _paint_paper(slide)
    _left_band(slide)
    head = _add_textbox(slide, Inches(0.6), Inches(0.45), Inches(12), Inches(0.7))
    hp = head.text_frame.paragraphs[0]
    hr = hp.add_run()
    _set_run(hr, title or "要点", size_pt=24, bold=True, color=NAVY)

    body = _add_textbox(slide, Inches(0.7), Inches(1.4), Inches(11.8), Inches(5.2))
    tf = body.text_frame
    tf.word_wrap = True
    items = [b.strip() for b in bullets if str(b).strip()][:8]
    if not items:
        items = ["（本页要点待补）"]
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(10)
        run = p.add_run()
        _set_run(run, f"•  {item}", size_pt=18, color=INK)
    _footer(slide, deck_title)


def render_closing(slide, title: str, bullets: list[str], deck_title: str) -> None:
    render_content(slide, title or "结语", bullets, deck_title)


def build(slides_path: Path, out_path: Path) -> dict:
    data = json.loads(slides_path.read_text(encoding="utf-8"))
    deck_title = str(data.get("title") or "汇报演示").strip()
    deck_sub = str(data.get("subtitle") or "").strip()
    slides = data.get("slides") or []
    if not isinstance(slides, list) or not slides:
        raise SystemExit("slides.json 缺少非空 slides 数组")

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    counts = {"title": 0, "section": 0, "content": 0, "closing": 0, "other": 0}
    for raw in slides:
        if not isinstance(raw, dict):
            continue
        stype = str(raw.get("type") or "content").lower().strip()
        title = str(raw.get("title") or "").strip()
        subtitle = str(raw.get("subtitle") or deck_sub).strip()
        bullets = raw.get("bullets") or []
        if not isinstance(bullets, list):
            bullets = [str(bullets)]

        slide = _blank_slide(prs)
        if stype == "title":
            render_title(slide, title or deck_title, subtitle, deck_title)
            counts["title"] += 1
        elif stype == "section":
            render_section(slide, title, deck_title)
            counts["section"] += 1
        elif stype == "closing":
            render_closing(slide, title, [str(x) for x in bullets], deck_title)
            counts["closing"] += 1
        else:
            render_content(slide, title, [str(x) for x in bullets], deck_title)
            counts["content"] += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return {
        "ok": True,
        "out": out_path.as_posix(),
        "pages": sum(counts.values()),
        "counts": counts,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build formal PPTX from slides JSON")
    parser.add_argument("--slides", required=True, help="Path to slides.json")
    parser.add_argument("--out", default="工作成果/汇报演示.pptx", help="Output pptx path")
    args = parser.parse_args(argv)

    cwd = Path.cwd()
    slides_path = confine_to_workspace(Path(args.slides), cwd)
    out_path = confine_to_workspace(Path(args.out), cwd)
    if not slides_path.is_file():
        print(f"找不到页纲: {slides_path}", file=sys.stderr)
        return 1
    try:
        result = build(slides_path, out_path)
    except json.JSONDecodeError as exc:
        print(f"slides.json 不是合法 JSON: {exc}", file=sys.stderr)
        return 1
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
