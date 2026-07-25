#!/usr/bin/env python3
"""按 Agent 标注的结构角色对 .docx 施加 GB/T 9704 样式。

Usage:
  python format_gongwen.py dump <docx> [--out paragraphs.json]
  python format_gongwen.py apply <docx> <roles.json> [--out formatted.docx]
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

# 西式二级：1.1 / 2.3.1 + 标题；半角括号二级：(一)
_RE_WESTERN_L2 = re.compile(
    r"^(?P<maj>\d+)\.(?P<min>\d+)(?:\.\d+)*"
    r"(?:[、.\s]+|\s*)(?P<title>\S.*)$"
)
_RE_WESTERN_L2_TIGHT = re.compile(
    r"^(?P<maj>\d+)\.(?P<min>\d+)(?:\.\d+)*(?P<title>[\u4e00-\u9fff].+)$"
)
_RE_HALFWIDTH_L2 = re.compile(
    r"^\((?P<num>[一二三四五六七八九十百零]+)\)\s*(?P<title>.*)$"
)
_RE_CN_L2 = re.compile(r"^（[一二三四五六七八九十百零]+）")
_RE_CN_L1 = re.compile(r"^[一二三四五六七八九十百零]+、")
_CN_DIGITS = "零一二三四五六七八九"

BODY_FONT_PT = 16
TITLE_FONT_PT = 22
COPY_FONT_PT = 14  # 四号，版记
LINE_SPACING_PT = 29
FIRST_LINE_INDENT_CHARS = 2
ASCII_FONT = "Times New Roman"

# role → (eastAsia font, size_pt, indent, align, color_rgb|None)
ROLE_STYLE: dict[str, dict[str, Any]] = {
    "masthead": {
        "font": "方正小标宋简体",
        "size": TITLE_FONT_PT,
        "indent": False,
        "align": WD_ALIGN_PARAGRAPH.CENTER,
        "color": RGBColor(0xFF, 0x00, 0x00),
    },
    "copy_no": {
        "font": "黑体",
        "size": BODY_FONT_PT,
        "indent": False,
        "align": WD_ALIGN_PARAGRAPH.LEFT,
        "color": None,
    },
    "secrecy": {
        "font": "黑体",
        "size": BODY_FONT_PT,
        "indent": False,
        "align": WD_ALIGN_PARAGRAPH.LEFT,
        "color": None,
    },
    "urgency": {
        "font": "黑体",
        "size": BODY_FONT_PT,
        "indent": False,
        "align": WD_ALIGN_PARAGRAPH.LEFT,
        "color": None,
    },
    "doc_number": {
        "font": "仿宋_GB2312",
        "size": BODY_FONT_PT,
        "indent": False,
        "align": WD_ALIGN_PARAGRAPH.CENTER,
        "color": None,
    },
    "signer": {
        "font": "仿宋_GB2312",
        "size": BODY_FONT_PT,
        "indent": False,
        "align": WD_ALIGN_PARAGRAPH.RIGHT,
        "color": None,
    },
    "title": {
        "font": "方正小标宋简体",
        "size": TITLE_FONT_PT,
        "indent": False,
        "align": WD_ALIGN_PARAGRAPH.CENTER,
        "color": None,
    },
    "main_recipient": {
        "font": "仿宋_GB2312",
        "size": BODY_FONT_PT,
        "indent": False,
        "align": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "color": None,
    },
    "h1": {
        "font": "黑体",
        "size": BODY_FONT_PT,
        "indent": True,
        "align": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "color": None,
    },
    "h2": {
        "font": "楷体_GB2312",
        "size": BODY_FONT_PT,
        "indent": True,
        "align": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "color": None,
    },
    "h3": {
        "font": "仿宋_GB2312",
        "size": BODY_FONT_PT,
        "indent": True,
        "align": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "color": None,
    },
    "h4": {
        "font": "仿宋_GB2312",
        "size": BODY_FONT_PT,
        "indent": True,
        "align": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "color": None,
    },
    "body": {
        "font": "仿宋_GB2312",
        "size": BODY_FONT_PT,
        "indent": True,
        "align": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "color": None,
    },
    "attachment_note": {
        "font": "仿宋_GB2312",
        "size": BODY_FONT_PT,
        "indent": True,
        "align": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "color": None,
    },
    "signature": {
        "font": "仿宋_GB2312",
        "size": BODY_FONT_PT,
        "indent": False,
        "align": WD_ALIGN_PARAGRAPH.RIGHT,
        "color": None,
    },
    "date": {
        "font": "仿宋_GB2312",
        "size": BODY_FONT_PT,
        "indent": False,
        "align": WD_ALIGN_PARAGRAPH.RIGHT,
        "color": None,
    },
    "annotation": {
        "font": "仿宋_GB2312",
        "size": BODY_FONT_PT,
        "indent": True,
        "align": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "color": None,
    },
    "annex_label": {
        "font": "黑体",
        "size": BODY_FONT_PT,
        "indent": False,
        "align": WD_ALIGN_PARAGRAPH.LEFT,
        "color": None,
    },
    "copy_to": {
        "font": "仿宋_GB2312",
        "size": COPY_FONT_PT,
        "indent": False,
        "align": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "color": None,
    },
    "print_info": {
        "font": "仿宋_GB2312",
        "size": COPY_FONT_PT,
        "indent": False,
        "align": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "color": None,
    },
    "empty": {
        "font": "仿宋_GB2312",
        "size": BODY_FONT_PT,
        "indent": False,
        "align": WD_ALIGN_PARAGRAPH.LEFT,
        "color": None,
    },
}

KNOWN_ROLES = set(ROLE_STYLE) | {"skip"}


def int_to_chinese(n: int) -> str:
    """将正整数转为中文数字（公文二级序数常用范围）。"""
    if n <= 0:
        raise ValueError(f"序数必须为正整数: {n}")
    if n < 10:
        return _CN_DIGITS[n]
    if n == 10:
        return "十"
    if n < 20:
        return "十" + _CN_DIGITS[n % 10]
    if n < 100:
        tens, ones = divmod(n, 10)
        return _CN_DIGITS[tens] + "十" + (_CN_DIGITS[ones] if ones else "")
    if n < 1000:
        hundreds, rest = divmod(n, 100)
        head = _CN_DIGITS[hundreds] + "百"
        if rest == 0:
            return head
        if rest < 10:
            return head + "零" + _CN_DIGITS[rest]
        return head + int_to_chinese(rest)
    return str(n)


def parse_western_l2(text: str) -> str | None:
    """若为西式多级标号（如 1.1 / 2.3.1），返回标题正文，否则 None。"""
    s = text.strip()
    m = _RE_WESTERN_L2.match(s)
    if m and m.group("title"):
        return m.group("title").strip()
    m = _RE_WESTERN_L2_TIGHT.match(s)
    if m:
        return m.group("title").strip()
    return None


def parse_halfwidth_l2(text: str) -> str | None:
    m = _RE_HALFWIDTH_L2.match(text.strip())
    if not m:
        return None
    return (m.group("title") or "").strip()


def looks_like_h1(text: str) -> bool:
    return bool(_RE_CN_L1.match(text.strip()))


def looks_like_cn_h2(text: str) -> bool:
    return bool(_RE_CN_L2.match(text.strip()))


def numbering_hint_for(text: str, h2_seq: int) -> dict[str, Any] | None:
    """dump 用：给出建议 role / 改正后 text（不改动计数器以外的状态）。"""
    s = text.strip()
    if not s:
        return None
    if looks_like_h1(s):
        return {"hint_role": "h1", "hint_text": None, "reset_h2": True}
    title = parse_western_l2(s)
    if title is not None:
        seq = h2_seq + 1
        return {
            "hint_role": "h2",
            "hint_text": f"（{int_to_chinese(seq)}）{title}",
            "bump_h2": True,
        }
    half = parse_halfwidth_l2(s)
    if half is not None:
        seq = h2_seq + 1
        body = half if half else s
        # 半角 (一)xxx → 全角（一）xxx；序数按出现顺序重排
        return {
            "hint_role": "h2",
            "hint_text": f"（{int_to_chinese(seq)}）{body}" if half else f"（{int_to_chinese(seq)}）",
            "bump_h2": True,
        }
    if looks_like_cn_h2(s):
        return {"hint_role": "h2", "hint_text": None, "bump_h2": True}
    return None


def resolve_numbering(
    paragraphs: list[Any],
    by_index: dict[int, dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[str]]:
    """按文档顺序解析最终 role 与 text；自动把 1.1 转为（一）。

    返回 (plan, warnings)。plan 项: index, role, text(可选), numbering_fixed(bool)
    """
    warnings: list[str] = []
    plan: list[dict[str, Any]] = []
    h2_seq = 0

    for i, para in enumerate(paragraphs):
        item = by_index.get(i)
        original = para.text
        original_stripped = original.strip()

        if item is None:
            role = "empty" if not original_stripped else "body"
            explicit_text = None
        else:
            role = str(item["role"])
            explicit_text = item.get("text")
            if explicit_text is not None:
                explicit_text = str(explicit_text)

        if role == "skip":
            plan.append({"index": i, "role": "skip", "text": None, "numbering_fixed": False})
            continue

        if role == "h1" or (role in ("body",) and looks_like_h1(original_stripped)):
            if role != "h1" and looks_like_h1(original_stripped):
                warnings.append(f"index={i} 疑似一级标题，已提升为 h1")
                role = "h1"
            h2_seq = 0

        western_title = parse_western_l2(original_stripped) if original_stripped else None
        half_title = parse_halfwidth_l2(original_stripped) if original_stripped else None
        needs_l2_fix = western_title is not None or half_title is not None

        if needs_l2_fix and role in ("body", "empty", "h3", "h4"):
            warnings.append(
                f"index={i} 西式/半角二级标号（原 role={role}），已提升为 h2 并改国标序数"
            )
            role = "h2"

        final_text = explicit_text
        numbering_fixed = False

        if role == "h2":
            h2_seq += 1
            marker = f"（{int_to_chinese(h2_seq)}）"
            if explicit_text is not None:
                # Agent 已给 text：若仍是西式则再改一次
                et = explicit_text.strip()
                et_western = parse_western_l2(et)
                et_half = parse_halfwidth_l2(et)
                if et_western is not None:
                    final_text = f"{marker}{et_western}"
                    numbering_fixed = True
                elif et_half is not None:
                    final_text = f"{marker}{et_half}" if et_half else marker
                    numbering_fixed = True
                elif not looks_like_cn_h2(et):
                    # 无序数的纯标题，补上（N）
                    final_text = f"{marker}{et}"
                    numbering_fixed = True
                else:
                    final_text = explicit_text
            elif western_title is not None:
                final_text = f"{marker}{western_title}"
                numbering_fixed = True
            elif half_title is not None:
                final_text = f"{marker}{half_title}" if half_title else marker
                numbering_fixed = True
            elif original_stripped and not looks_like_cn_h2(original_stripped):
                # 标成 h2 但无国标序数
                final_text = f"{marker}{original_stripped}"
                numbering_fixed = True
                warnings.append(f"index={i} h2 缺少国标序数，已补 {marker}")

        plan.append(
            {
                "index": i,
                "role": role,
                "text": final_text,
                "numbering_fixed": numbering_fixed,
            }
        )

    return plan, warnings


def set_page_margins(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(37)
        section.bottom_margin = Mm(35)
        section.left_margin = Mm(28)
        section.right_margin = Mm(26)


def set_default_font(
    doc: Document,
    font_name: str = "仿宋_GB2312",
    font_name_ascii: str = ASCII_FONT,
    size: Pt = Pt(BODY_FONT_PT),
) -> None:
    style = doc.styles["Normal"]
    font = style.font
    font.name = font_name
    font.size = size
    r = style.element.rPr
    if r is None:
        r = OxmlElement("w:rPr")
        style.element.append(r)
    rFonts = r.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        r.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), font_name)
    rFonts.set(qn("w:ascii"), font_name_ascii)
    rFonts.set(qn("w:hAnsi"), font_name_ascii)


def set_line_spacing(doc: Document, line_spacing: float = LINE_SPACING_PT) -> None:
    style = doc.styles["Normal"]
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(line_spacing)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)


def set_first_line_indent_chars(para, chars: int = FIRST_LINE_INDENT_CHARS) -> None:
    p_element = para._element
    pPr = p_element.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        p_element.insert(0, pPr)
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    for attr in ("firstLine", "hanging"):
        key = qn(f"w:{attr}")
        if ind.get(key) is not None:
            del ind.attrib[key]
    ind.set(qn("w:firstLineChars"), str(int(chars * 100)))
    ind.set(qn("w:firstLine"), str(int(BODY_FONT_PT * chars * 20)))


def clear_first_line_indent(para) -> None:
    para.paragraph_format.first_line_indent = Pt(0)
    p_element = para._element
    pPr = p_element.find(qn("w:pPr"))
    if pPr is None:
        return
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        return
    for attr in ("firstLine", "firstLineChars", "hanging"):
        key = qn(f"w:{attr}")
        if ind.get(key) is not None:
            del ind.attrib[key]


def apply_paragraph_format(para, *, indent: bool, align) -> None:
    para.alignment = align
    pf = para.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(LINE_SPACING_PT)
    if indent:
        set_first_line_indent_chars(para, FIRST_LINE_INDENT_CHARS)
    else:
        clear_first_line_indent(para)


def _set_run_font(run, font_name: str, size: Pt, bold: bool = False, color=None) -> None:
    run.font.name = font_name
    run.font.size = size
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    r = run._element.rPr
    if r is None:
        r = OxmlElement("w:rPr")
        run._element.append(r)
    rFonts = r.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        r.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), font_name)
    rFonts.set(qn("w:ascii"), ASCII_FONT)
    rFonts.set(qn("w:hAnsi"), ASCII_FONT)


def adjust_paragraph_font(para, font_name: str, size: Pt, bold: bool = False, color=None) -> None:
    for run in para.runs:
        _set_run_font(run, font_name, size, bold=bold, color=color)


def replace_paragraph_text(para, text: str, font_name: str, size: Pt, color=None) -> None:
    """整段替换文本（用于标号规范化），单 run 施加字体。"""
    # 清空现有 runs
    for run in list(para.runs):
        run._element.getparent().remove(run._element)
    run = para.add_run(text)
    _set_run_font(run, font_name, size, color=color)


def insert_red_line(para, mm_below: float = 4) -> None:
    p_element = para._element
    pPr = p_element.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        p_element.insert(0, pPr)
    # 移除旧底边框，避免重复叠加
    old = pPr.find(qn("w:pBdr"))
    if old is not None:
        pPr.remove(old)
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), str(int(mm_below * 56.7)))
    bottom.set(qn("w:color"), "FF0000")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _alignment_name(para) -> str | None:
    align = para.alignment
    if align is None:
        return None
    mapping = {
        WD_ALIGN_PARAGRAPH.LEFT: "left",
        WD_ALIGN_PARAGRAPH.CENTER: "center",
        WD_ALIGN_PARAGRAPH.RIGHT: "right",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
    }
    return mapping.get(align, str(align))


def cmd_dump(docx_path: Path, out_path: Path | None) -> dict[str, Any]:
    doc = Document(str(docx_path))
    paragraphs = []
    h2_seq = 0
    western_l2_count = 0
    for i, para in enumerate(doc.paragraphs):
        text = para.text
        stripped = text.strip()
        entry: dict[str, Any] = {
            "index": i,
            "text": text,
            "stripped": stripped,
            "empty": not stripped,
            "alignment": _alignment_name(para),
            "style": para.style.name if para.style is not None else None,
        }
        hint = numbering_hint_for(stripped, h2_seq) if stripped else None
        if hint:
            if hint.get("reset_h2"):
                h2_seq = 0
            entry["hint_role"] = hint["hint_role"]
            if hint.get("hint_text"):
                entry["hint_text"] = hint["hint_text"]
            if hint.get("bump_h2"):
                h2_seq += 1
            if parse_western_l2(stripped) is not None or parse_halfwidth_l2(stripped) is not None:
                western_l2_count += 1
                entry["needs_numbering_fix"] = True
        paragraphs.append(entry)
    payload = {
        "source": str(docx_path),
        "paragraph_count": len(paragraphs),
        "paragraphs": paragraphs,
        "roles_schema_hint": {
            "required_fields": ["index", "role"],
            "optional_fields": ["text"],
            "roles": sorted(KNOWN_ROLES),
            "note": (
                "空段用 empty；不确定用 skip。"
                "西式二级（1.1）请标 h2；apply 会自动改为（一）（二），"
                "亦可直接采用 hint_text。"
            ),
        },
    }
    if out_path is None:
        out_path = docx_path.with_name(docx_path.stem + "_paragraphs.json")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return {
        "ok": True,
        "action": "dump",
        "source": str(docx_path),
        "paragraphs_json": str(out_path),
        "paragraph_count": len(paragraphs),
        "non_empty": sum(1 for p in paragraphs if not p["empty"]),
        "western_l2_candidates": western_l2_count,
    }


def _load_roles(roles_path: Path) -> dict[str, Any]:
    data = json.loads(roles_path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("roles.json 根节点必须是对象")
    paras = data.get("paragraphs")
    if not isinstance(paras, list) or not paras:
        raise ValueError("roles.json 必须包含非空 paragraphs 数组")
    return data


def cmd_apply(docx_path: Path, roles_path: Path, out_path: Path | None) -> dict[str, Any]:
    roles_data = _load_roles(roles_path)
    doc = Document(str(docx_path))
    set_page_margins(doc)
    set_default_font(doc)
    set_line_spacing(doc)

    n = len(doc.paragraphs)
    by_index: dict[int, dict[str, Any]] = {}
    errors: list[str] = []
    warnings: list[str] = []

    for item in roles_data["paragraphs"]:
        if not isinstance(item, dict):
            errors.append(f"非法条目: {item!r}")
            continue
        if "index" not in item or "role" not in item:
            errors.append(f"缺少 index/role: {item!r}")
            continue
        idx = int(item["index"])
        role = str(item["role"]).strip()
        if role not in KNOWN_ROLES:
            errors.append(f"未知 role={role!r} @ index={idx}")
            continue
        if idx < 0 or idx >= n:
            errors.append(f"index={idx} 越界（共 {n} 段）")
            continue
        if idx in by_index:
            warnings.append(f"index={idx} 重复标注，后写覆盖先写")
        by_index[idx] = item

    if errors:
        raise ValueError("; ".join(errors))

    unassigned = [i for i in range(n) if i not in by_index]
    if unassigned:
        warnings.append(
            f"{len(unassigned)} 段未标注，将按 empty/body 兜底: "
            + ",".join(str(i) for i in unassigned[:20])
            + ("..." if len(unassigned) > 20 else "")
        )

    plan, num_warnings = resolve_numbering(list(doc.paragraphs), by_index)
    warnings.extend(num_warnings)

    applied: dict[str, int] = {}
    text_rewrites = 0
    numbering_fixed = 0

    for step in plan:
        i = step["index"]
        para = doc.paragraphs[i]
        role = step["role"]

        if role == "skip":
            applied["skip"] = applied.get("skip", 0) + 1
            continue

        style = ROLE_STYLE[role]
        apply_paragraph_format(para, indent=bool(style["indent"]), align=style["align"])

        new_text = step.get("text")
        if new_text is not None:
            replace_paragraph_text(
                para,
                str(new_text),
                style["font"],
                Pt(style["size"]),
                color=style["color"],
            )
            text_rewrites += 1
            if step.get("numbering_fixed"):
                numbering_fixed += 1
        else:
            adjust_paragraph_font(
                para,
                style["font"],
                Pt(style["size"]),
                color=style["color"],
            )
            if not para.runs and role != "empty":
                run = para.add_run(para.text)
                _set_run_font(run, style["font"], Pt(style["size"]), color=style["color"])

        applied[role] = applied.get(role, 0) + 1

    red_after = roles_data.get("red_line_after")
    red_line = False
    if red_after is not None:
        idx = int(red_after)
        if 0 <= idx < n:
            insert_red_line(doc.paragraphs[idx])
            red_line = True
        else:
            warnings.append(f"red_line_after={idx} 越界，已忽略")

    if out_path is None:
        out_path = docx_path.with_name(docx_path.stem + "_formatted.docx")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))

    return {
        "ok": True,
        "action": "apply",
        "source": str(docx_path),
        "roles_json": str(roles_path),
        "output": str(out_path),
        "paragraph_count": n,
        "applied_roles": applied,
        "text_rewrites": text_rewrites,
        "numbering_fixed": numbering_fixed,
        "red_line": red_line,
        "warnings": warnings,
        "notes": roles_data.get("notes") or [],
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="公文排版：dump 段落供 Agent 标注，apply 按角色施加样式"
    )
    sub = parser.add_subparsers(dest="command", required=True)

    p_dump = sub.add_parser("dump", help="导出段落列表 JSON")
    p_dump.add_argument("docx", type=Path)
    p_dump.add_argument("--out", type=Path, default=None)

    p_apply = sub.add_parser("apply", help="按 roles.json 施加样式")
    p_apply.add_argument("docx", type=Path)
    p_apply.add_argument("roles", type=Path)
    p_apply.add_argument("--out", type=Path, default=None)

    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)

    try:
        if args.command == "dump":
            if not args.docx.is_file():
                raise FileNotFoundError(f"文件不存在: {args.docx}")
            summary = cmd_dump(args.docx, args.out)
        elif args.command == "apply":
            if not args.docx.is_file():
                raise FileNotFoundError(f"文件不存在: {args.docx}")
            if not args.roles.is_file():
                raise FileNotFoundError(f"roles 不存在: {args.roles}")
            summary = cmd_apply(args.docx, args.roles, args.out)
        else:
            raise ValueError(f"未知命令: {args.command}")
    except Exception as exc:  # noqa: BLE001 — CLI 边界统一报错
        print(f"ERROR: {exc}", file=sys.stderr)
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False))
        return 1

    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
