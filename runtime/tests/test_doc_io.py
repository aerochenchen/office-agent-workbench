"""Tests for document normalize + workspace_extract."""

from __future__ import annotations

from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document
from openpyxl import Workbook

from office_agent.audit import AuditLog
from office_agent.doc_io import DocIOError, extract_file, extract_pdf, normalize_path
from office_agent.skills import SkillRegistry
from office_agent.tools import ToolExecutor
from office_agent.workspace import Workspace


@pytest.fixture(autouse=True)
def _stub_pdf_ocr_page_text(monkeypatch: pytest.MonkeyPatch) -> None:
    """Keep PDF unit tests off the real RapidOCR engine unless a test opts in."""
    monkeypatch.setattr(
        "office_agent.doc_io._pdf_ocr_page_text",
        lambda *_a, **_k: "",
        raising=False,
    )


def _write_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("专项汇报", level=1)
    doc.add_paragraph("一、工作进展")
    doc.add_paragraph("建成智能平台，覆盖基层网点 128 个。")
    doc.add_paragraph("二、下一步打算")
    doc.add_paragraph("持续推进重点任务落地。")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def _write_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"
    ws.append(["部门", "完成率", "备注"])
    ws.append(["办公室", "95%", "含试点"])
    ws.append(["业务处", "88%", ""])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


def _write_xls(path: Path) -> None:
    xlwt = pytest.importorskip("xlwt")
    book = xlwt.Workbook()
    sheet = book.add_sheet("数据")
    sheet.write(0, 0, "项目")
    sheet.write(0, 1, "金额")
    sheet.write(1, 0, "数字化")
    sheet.write(1, 1, 128)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(str(path))


def _write_pdf_pages(path: Path, page_texts: list[str | None]) -> None:
    """Write a minimal multi-page PDF. None = empty content (no text layer)."""
    pytest.importorskip("reportlab")
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    path.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(path), pagesize=A4)
    for text in page_texts:
        if text:
            c.setFont("Helvetica", 12)
            y = 800
            for line in text.split("\n"):
                c.drawString(72, y, line)
                y -= 18
        c.showPage()
    c.save()


def _write_pdf_table_lines_only(path: Path) -> None:
    """PDF page with drawn grid lines but no text layer (triggers find_tables)."""
    pytest.importorskip("reportlab")
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    path.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(path), pagesize=A4)
    x0, y0 = 72, 600
    w, h = 300, 100
    rows, cols = 3, 3
    cell_w = w / cols
    cell_h = h / rows
    for i in range(rows + 1):
        y = y0 - i * cell_h
        c.line(x0, y, x0 + w, y)
    for j in range(cols + 1):
        x = x0 + j * cell_w
        c.line(x, y0, x, y0 - h)
    c.showPage()
    c.save()


def _write_pdf_with_table(path: Path) -> None:
    pytest.importorskip("reportlab")
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(path), pagesize=A4)
    styles = getSampleStyleSheet()
    data = [
        ["Dept", "Rate", "Note"],
        ["Office", "95%", "pilot"],
        ["Biz", "88%", ""],
    ]
    table = Table(data, colWidths=[120, 80, 120])
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ]
        )
    )
    doc.build(
        [
            Paragraph("Progress report", styles["Heading1"]),
            Spacer(1, 12),
            table,
        ]
    )


def test_extract_docx_units(tmp_path: Path):
    src = tmp_path / "a.docx"
    _write_docx(src)
    result = extract_file(src, tmp_path)
    assert result.ok is True
    assert result.format == "docx"
    assert result.doc_key
    assert result.units
    blob = "\n".join(u.text for u in result.units)
    assert "128" in blob
    assert result.units[0].unit_id.startswith(result.doc_key + "#u")
    assert result.units[0].meta.get("granularity") == "section"
    assert "para_start" in result.units[0].meta
    d = result.units[0].to_dict()
    assert "meta" in d


def test_extract_docx_paragraph_granularity(tmp_path: Path):
    src = tmp_path / "para.docx"
    _write_docx(src)
    section = extract_file(src, tmp_path, granularity="section")
    paragraph = extract_file(src, tmp_path, granularity="paragraph")
    assert paragraph.ok is True
    # _write_docx: heading + 4 body paras = 5 non-empty paragraphs
    para_units = [u for u in paragraph.units if u.kind in {"paragraph", "heading"}]
    assert len(para_units) == 5
    assert len(para_units) >= len(section.units)
    assert all(u.meta.get("granularity") == "paragraph" for u in para_units)
    assert all(u.meta.get("para_start") == u.meta.get("para_end") for u in para_units)
    assert any(u.kind == "heading" for u in para_units)


def test_extract_xlsx_units(tmp_path: Path):
    src = tmp_path / "b.xlsx"
    _write_xlsx(src)
    result = extract_file(src, tmp_path)
    assert result.ok is True
    assert result.format == "xlsx"
    blob = "\n".join(u.text for u in result.units)
    assert "办公室" in blob
    assert "95%" in blob


def test_extract_xlsx_cells_granularity(tmp_path: Path):
    src = tmp_path / "cells.xlsx"
    _write_xlsx(src)
    result = extract_file(src, tmp_path, granularity="cells")
    assert result.ok is True
    cells = [u for u in result.units if u.kind == "cell"]
    assert len(cells) >= 6  # header 3 + data 3+
    assert any(u.locator.endswith("!B2") or "B2" in u.locator for u in cells)
    assert any("95%" in u.text for u in cells)
    assert all(u.meta.get("granularity") == "cells" for u in cells)
    assert any(u.meta.get("header") == "完成率" for u in cells if u.meta.get("row", 0) > 1)


def test_normalize_xls_to_xlsx(tmp_path: Path):
    src = tmp_path / "legacy.xls"
    _write_xls(src)
    modern, fmt, _ = normalize_path(src, tmp_path)
    assert fmt == "xlsx"
    assert modern.suffix.lower() == ".xlsx"
    assert modern.is_file()
    result = extract_file(src, tmp_path)
    assert result.ok is True
    assert result.normalized_path
    blob = "\n".join(u.text for u in result.units)
    assert "数字化" in blob


def test_normalize_doc_uses_converter(tmp_path: Path):
    src = tmp_path / "old.doc"
    src.write_bytes(b"fake-doc-bytes")
    dest = tmp_path / ".office-agent" / "work" / "normalized" / "old.docx"

    def fake_soffice(cmd, **kwargs):
        _write_docx(dest)

        class R:
            returncode = 0
            stdout = ""
            stderr = ""

        return R()

    with patch("office_agent.doc_io._which", return_value="/usr/bin/soffice"):
        with patch("office_agent.doc_io.subprocess.run", side_effect=fake_soffice):
            modern, fmt, _ = normalize_path(src, tmp_path, force=True)
    assert fmt == "docx"
    assert modern == dest
    result = extract_file(src, tmp_path)
    assert result.ok is True
    assert any("128" in u.text for u in result.units)


def test_normalize_doc_fails_clearly_without_converter(tmp_path: Path):
    src = tmp_path / "old.doc"
    src.write_bytes(b"fake")
    with patch("office_agent.doc_io._which", return_value=None):
        with patch("office_agent.doc_io.shutil.which", return_value=None):
            with patch("office_agent.doc_io._find_soffice", return_value=None):
                with patch(
                    "office_agent.doc_io._normalize_doc_win32", return_value=False
                ):
                    with pytest.raises(DocIOError, match="无法将 .doc"):
                        normalize_path(src, tmp_path, force=True)


def test_find_soffice_checks_windows_install_paths(tmp_path: Path, monkeypatch):
    fake = tmp_path / "LibreOffice" / "program" / "soffice.exe"
    fake.parent.mkdir(parents=True)
    fake.write_bytes(b"x")
    monkeypatch.setattr("office_agent.doc_io.sys.platform", "win32")
    monkeypatch.setenv("PROGRAMFILES", str(tmp_path))
    monkeypatch.delenv("PROGRAMFILES(X86)", raising=False)
    monkeypatch.delenv("LOCALAPPDATA", raising=False)
    with patch("office_agent.doc_io._which", return_value=None):
        from office_agent.doc_io import _find_soffice

        assert _find_soffice() == str(fake)


def test_normalize_doc_uses_windows_word_com(tmp_path: Path, monkeypatch):
    src = tmp_path / "old.doc"
    src.write_bytes(b"fake-doc-bytes")
    dest = tmp_path / ".office-agent" / "work" / "normalized" / "old.docx"

    def fake_win32(src_path: Path, dest_path: Path) -> bool:
        _write_docx(dest_path)
        return True

    monkeypatch.setattr("office_agent.doc_io.sys.platform", "win32")
    with patch("office_agent.doc_io._find_soffice", return_value=None):
        with patch("office_agent.doc_io._normalize_doc_win32", side_effect=fake_win32):
            modern, fmt, _ = normalize_path(src, tmp_path, force=True)
    assert fmt == "docx"
    assert modern == dest
    assert dest.is_file()


def test_workspace_extract_tool(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("OFFICE_AGENT_DATA", str(tmp_path / "data"))
    (tmp_path / "data" / "skills").mkdir(parents=True)
    ws = tmp_path / "ws"
    ws.mkdir()
    _write_docx(ws / "汇报.docx")
    _write_xlsx(ws / "表.xlsx")
    ex = ToolExecutor(
        Workspace(ws),
        SkillRegistry(),
        permission_mode="trust",
        audit=AuditLog(tmp_path / "data" / "a.sqlite"),
    )
    r1 = ex.execute("workspace_extract", {"path": "汇报.docx"})
    assert r1["ok"] is True
    assert r1["format"] == "docx"
    assert r1["units"]
    r2 = ex.execute("workspace_extract", {"path": "表.xlsx"})
    assert r2["ok"] is True
    assert any("办公室" in u["text"] for u in r2["units"])


def test_workspace_extract_xls_via_tool(tmp_path: Path, monkeypatch):
    pytest.importorskip("xlwt")
    monkeypatch.setenv("OFFICE_AGENT_DATA", str(tmp_path / "data"))
    (tmp_path / "data" / "skills").mkdir(parents=True)
    ws = tmp_path / "ws"
    ws.mkdir()
    _write_xls(ws / "旧表.xls")
    ex = ToolExecutor(Workspace(ws), SkillRegistry(), permission_mode="trust")
    r = ex.execute("workspace_extract", {"path": "旧表.xls"})
    assert r["ok"] is True
    assert r["format"] == "xlsx"
    assert r.get("normalized_path")
    assert any("数字化" in u["text"] for u in r["units"])


def test_workspace_extract_passes_granularity(tmp_path: Path):
    (tmp_path / "data" / "skills").mkdir(parents=True)
    ws = tmp_path / "ws"
    ws.mkdir()
    _write_docx(ws / "稿.docx")
    ex = ToolExecutor(Workspace(ws), SkillRegistry(), permission_mode="trust")
    r = ex.execute(
        "workspace_extract",
        {"path": "稿.docx", "granularity": "paragraph"},
    )
    assert r["ok"] is True
    assert len([u for u in r["units"] if u["kind"] in {"paragraph", "heading"}]) == 5
    assert r["units"][0]["meta"]["granularity"] == "paragraph"


def test_extract_pdf_pages_and_empty_warning(tmp_path: Path):
    src = tmp_path / "mixed.pdf"
    _write_pdf_pages(src, ["Alpha progress 128 sites", None])
    result = extract_file(src, tmp_path)
    assert result.ok is True
    assert result.format == "pdf"
    assert result.doc_key
    page_units = [u for u in result.units if u.kind == "page"]
    assert len(page_units) >= 1
    blob = "\n".join(u.text for u in page_units)
    assert "128" in blob
    empty_page_warnings = [
        w for w in result.warnings if "无文本层" in w and "第2页" in w
    ]
    assert empty_page_warnings, "empty page must warn even without table detection"
    assert not any(u.kind == "table" for u in result.units)


def test_extract_pdf_max_units_sets_truncated(tmp_path: Path):
    src = tmp_path / "many_pages.pdf"
    _write_pdf_pages(src, ["Page one", "Page two", "Page three"])
    result = extract_file(src, tmp_path, max_units=1)
    assert result.ok is True
    assert len(result.units) == 1
    assert result.truncated is True


def test_extract_pdf_aggregates_empty_page_warnings(tmp_path: Path):
    src = tmp_path / "empty_pages.pdf"
    _write_pdf_pages(src, [None] * 12)
    result = extract_file(src, tmp_path)
    no_text_warnings = [w for w in result.warnings if "无文本层" in w]
    assert len(no_text_warnings) == 1
    assert "共12页" in no_text_warnings[0]


def test_extract_pdf_table_lines_no_text_warns(tmp_path: Path):
    """Grid-only page: find_tables may fire but chars are empty — must still warn."""
    src = tmp_path / "grid_only.pdf"
    _write_pdf_table_lines_only(src)
    result = extract_file(src, tmp_path)
    assert result.ok is True
    assert any("无文本层" in w for w in result.warnings)
    assert not any(u.kind == "table" for u in result.units)
    assert not any(u.text.strip("| \t") for u in result.units)


def test_extract_pdf_skips_empty_cell_table_rows(tmp_path: Path, monkeypatch):
    """Empty-cell grids must not produce pseudo table units (pipe-only rows)."""
    import pdfplumber

    class _FakeTable:
        bbox = (0.0, 0.0, 100.0, 100.0)

        def extract(self):
            return [["", "", ""], ["", "", ""]]

    class _FakePage:
        chars: list[dict[str, str]] = []
        cache_flushed = False

        def find_tables(self):
            return [_FakeTable()]

        def filter(self, _pred):
            return self

        def extract_text(self):
            return ""

        def flush_cache(self):
            self.cache_flushed = True

    class _FakePdf:
        pages = [_FakePage()]

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

    monkeypatch.setattr(pdfplumber, "open", lambda _path: _FakePdf())
    units, _truncated, warnings = extract_pdf(
        tmp_path / "fake.pdf", "abc123", max_chars=8000, max_units=50
    )
    assert not units
    assert any("无文本层" in w for w in warnings)
    assert _FakePdf.pages[0].cache_flushed is True


def test_extract_pdf_falls_back_when_table_filter_loses_page_text(
    tmp_path: Path, monkeypatch
):
    import pdfplumber

    class _FakeTable:
        bbox = (0.0, 0.0, 100.0, 100.0)

        def extract(self):
            return [["", ""]]

    class _FilteredPage:
        def extract_text(self):
            return ""

    class _FakePage:
        chars = [{"text": "Recovered"}]

        def find_tables(self):
            return [_FakeTable()]

        def filter(self, _pred):
            return _FilteredPage()

        def extract_text(self):
            return "Recovered body"

        def flush_cache(self):
            pass

    class _FakePdf:
        pages = [_FakePage()]

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

    monkeypatch.setattr(pdfplumber, "open", lambda _path: _FakePdf())
    units, truncated, _warnings = extract_pdf(
        tmp_path / "fake.pdf", "abc123", max_chars=8000, max_units=50
    )
    assert truncated is False
    assert any(unit.text == "Recovered body" for unit in units)


def test_extract_pdf_table_unit(tmp_path: Path):
    src = tmp_path / "table.pdf"
    _write_pdf_with_table(src)
    result = extract_file(src, tmp_path)
    assert result.ok is True
    tables = [u for u in result.units if u.kind == "table"]
    assert tables, "expected at least one table unit"
    assert any("|" in u.text for u in tables)
    assert any("Office" in u.text or "95%" in u.text for u in tables)


def test_extract_pdf_paragraph_granularity(tmp_path: Path):
    src = tmp_path / "para.pdf"
    _write_pdf_pages(src, ["Line one\nLine two"])
    section = extract_file(src, tmp_path, granularity="section")
    paragraph = extract_file(src, tmp_path, granularity="paragraph")
    assert paragraph.ok is True
    para_units = [u for u in paragraph.units if u.kind == "paragraph"]
    assert len(para_units) >= 1
    assert all(u.meta.get("granularity") == "paragraph" for u in para_units)
    assert section.units[0].kind == "page"


def test_extract_pdf_cells_falls_back(tmp_path: Path):
    src = tmp_path / "cells.pdf"
    _write_pdf_pages(src, ["fallback body"])
    result = extract_file(src, tmp_path, granularity="cells")
    assert result.ok is True
    assert result.format == "pdf"
    assert any(u.kind == "page" for u in result.units)


def test_workspace_extract_pdf_tool(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("OFFICE_AGENT_DATA", str(tmp_path / "data"))
    (tmp_path / "data" / "skills").mkdir(parents=True)
    ws = tmp_path / "ws"
    ws.mkdir()
    _write_pdf_pages(ws / "材料.pdf", ["Workspace PDF body 42"])
    ex = ToolExecutor(
        Workspace(ws),
        SkillRegistry(),
        permission_mode="trust",
        audit=AuditLog(tmp_path / "data" / "a.sqlite"),
    )
    r = ex.execute("workspace_extract", {"path": "材料.pdf"})
    assert r["ok"] is True
    assert r["format"] == "pdf"
    assert r["units"]
    assert r["units"][0]["unit_id"].startswith(r["doc_key"] + "#u")
    assert any("42" in u["text"] for u in r["units"])


def test_extract_pdf_uses_body_even_if_chars_empty(tmp_path: Path, monkeypatch):
    """pdfminer chars may be empty on some CJK PDFs while extract_text still works."""
    import pdfplumber

    class _FakePage:
        chars: list[dict[str, str]] = []

        def find_tables(self):
            return []

        def extract_text(self):
            return "正文仍可读 128"

        def flush_cache(self):
            pass

    class _FakePdf:
        pages = [_FakePage()]

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

    monkeypatch.setattr(pdfplumber, "open", lambda _path: _FakePdf())
    units, _truncated, warnings = extract_pdf(
        tmp_path / "fake.pdf", "abc123", max_chars=8000, max_units=50
    )
    assert any("128" in u.text for u in units)
    assert not any("无文本层" in w for w in warnings)


def test_extract_pdf_falls_back_to_pdfium_when_pdfminer_empty(
    tmp_path: Path, monkeypatch
):
    """When pdfminer sees no text layer, try pypdfium2 before calling it a scan."""
    import pdfplumber

    src = tmp_path / "fallback.pdf"
    _write_pdf_pages(src, ["PDFium recovered text 99"])

    class _FakePage:
        chars: list[dict[str, str]] = []

        def find_tables(self):
            return []

        def extract_text(self):
            return ""

        def flush_cache(self):
            pass

    class _FakePdf:
        pages = [_FakePage()]

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

    monkeypatch.setattr(pdfplumber, "open", lambda _path: _FakePdf())
    result = extract_file(src, tmp_path)
    assert result.ok is True
    assert any("99" in u.text for u in result.units)
    assert not any("无文本层" in w for w in result.warnings)


def test_extract_pdf_ocr_fills_empty_scan_page(tmp_path: Path, monkeypatch):
    src = tmp_path / "scan.pdf"
    _write_pdf_pages(src, [None])
    monkeypatch.setattr(
        "office_agent.doc_io._pdf_ocr_page_text",
        lambda *_a, **_k: "扫描通知 128 个网点",
    )
    result = extract_file(src, tmp_path)
    assert result.ok is True
    pages = [u for u in result.units if u.kind == "page"]
    assert any("128" in u.text for u in pages)
    assert all(u.meta.get("source") == "ocr" for u in pages)
    assert any("印刷体识别" in w for w in result.warnings)
    assert any("处理得好的材料" in w and "不能保证准确" in w for w in result.warnings)
    assert not any("请勿当原文核对" in w for w in result.warnings)
    assert not any("暂不支持 OCR" in w for w in result.warnings)


def test_extract_pdf_does_not_ocr_text_layer_pages(tmp_path: Path, monkeypatch):
    src = tmp_path / "digital.pdf"
    _write_pdf_pages(src, ["Digital body 42"])
    called = {"n": 0}

    def _should_not_run(*_a, **_k) -> str:
        called["n"] += 1
        return "SHOULD_NOT"

    monkeypatch.setattr("office_agent.doc_io._pdf_ocr_page_text", _should_not_run)
    result = extract_file(src, tmp_path)
    assert called["n"] == 0
    assert any("42" in u.text for u in result.units)
    assert all(u.meta.get("source") != "ocr" for u in result.units)


def test_extract_pdf_ocr_mixed_digital_and_scan(tmp_path: Path, monkeypatch):
    src = tmp_path / "mixed.pdf"
    _write_pdf_pages(src, ["Alpha progress 128 sites", None])
    monkeypatch.setattr(
        "office_agent.doc_io._pdf_ocr_page_text",
        lambda *_a, **_k: "封面扫描件",
    )
    result = extract_file(src, tmp_path)
    pages = [u for u in result.units if u.kind == "page"]
    digital = [u for u in pages if "128" in u.text]
    ocr = [u for u in pages if u.meta.get("source") == "ocr"]
    assert digital
    assert any("封面扫描件" in u.text for u in ocr)
    assert all(u.meta.get("source") != "ocr" for u in digital)


def test_extract_pdf_ocr_max_pages_sets_truncated(tmp_path: Path, monkeypatch):
    src = tmp_path / "many_scans.pdf"
    _write_pdf_pages(src, [None, None, None])
    monkeypatch.setattr(
        "office_agent.doc_io._pdf_ocr_page_text",
        lambda *_a, **_k: "scan line",
    )
    result = extract_file(src, tmp_path, max_ocr_pages=1)
    ocr_units = [u for u in result.units if u.meta.get("source") == "ocr"]
    assert len(ocr_units) == 1
    assert result.truncated is True
    assert any("上限" in w for w in result.warnings)
