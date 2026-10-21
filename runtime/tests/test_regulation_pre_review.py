"""制度发布前审查技能：安装校验通过，脚本只出 xlsx 和底稿。"""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from office_agent.skill_validate import validate_skill_dir

SKILL_DIR = Path(__file__).resolve().parents[2] / "bundled" / "skills" / "regulation-pre-review"
SAMPLE = SKILL_DIR / "assets" / "sample" / "sample_review_data.json"


def _run(args: list[str]) -> subprocess.CompletedProcess[str]:
    """Match Runtime script capture: locale text, not a forced UTF-8 pipe."""
    return subprocess.run(args, capture_output=True, text=True)


def test_skill_package_passes_install_gate():
    result = validate_skill_dir(SKILL_DIR)
    assert result["ok"] is True, result["errors"]
    warnings = "\n".join(result["warnings"])
    assert "description is" not in warnings
    assert "何时" not in warnings
    assert "步数预算" not in warnings
    blob = "\n".join(p.read_text(encoding="utf-8") for p in (SKILL_DIR / "scripts").glob("*.py"))
    assert "pymupdf" not in blob
    assert "win32com" not in blob
    assert "fitz" not in blob


def test_build_report_writes_xlsx_and_draft(tmp_path: Path):
    out = tmp_path / "审查报告.xlsx"
    draft = tmp_path / "审查底稿.md"
    proc = _run(
        [
            sys.executable,
            str(SKILL_DIR / "scripts" / "build_report.py"),
            str(SAMPLE),
            "-o",
            str(out),
            "--draft",
            str(draft),
        ]
    )
    assert proc.returncode == 0, proc.stderr
    summary = json.loads(proc.stdout)
    assert summary["ok"] is True
    assert summary["items"] == 22
    assert Path(summary["xlsx"]).is_file()
    assert Path(summary["draft"]).is_file()

    from office_agent.xlsx_check import xlsx_package_errors

    assert xlsx_package_errors(out, "审查报告.xlsx") == []

    wb = load_workbook(out)
    ws = wb.active
    assert ws.title == "规范制度发布前审查报告"
    assert ws["A2"].value == "规范制度发布前审查报告"
    assert str(ws["D9"].value).startswith("符合")
    assert ws["A36"].value and "审查不通过" in str(ws["A36"].value)
    text = draft.read_text(encoding="utf-8")
    assert "待补材料" in text
    assert "审查不通过" in text


def test_build_report_rejects_short_item_list(tmp_path: Path):
    data = json.loads(SAMPLE.read_text(encoding="utf-8"))
    data["items"] = data["items"][:3]
    path = tmp_path / "short.json"
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    proc = _run(
        [
            sys.executable,
            str(SKILL_DIR / "scripts" / "build_report.py"),
            str(path),
            "-o",
            str(tmp_path / "out.xlsx"),
        ]
    )
    assert proc.returncode != 0
    assert "ERROR:" in proc.stderr
    assert "22" in proc.stderr


def test_build_report_rejects_pdf_flag(tmp_path: Path):
    proc = _run(
        [
            sys.executable,
            str(SKILL_DIR / "scripts" / "build_report.py"),
            str(SAMPLE),
            "-o",
            str(tmp_path / "out.xlsx"),
            "--pdf",
        ]
    )
    assert proc.returncode != 0
    assert not (tmp_path / "out.xlsx").exists()


def test_extract_docx_keeps_paragraph_text(tmp_path: Path):
    src = tmp_path / "制度.docx"
    doc = Document()
    doc.add_paragraph("第一条 为规范无人驾驶航空器干扰事件处置，制定本办法。")
    doc.add_paragraph("")
    doc.add_paragraph("第七条 本站对违法放飞处以罚款。")
    doc.save(src)
    out = tmp_path / "取稿.txt"
    proc = _run(
        [
            sys.executable,
            str(SKILL_DIR / "scripts" / "extract_docx.py"),
            str(src),
            "-o",
            str(out),
        ]
    )
    assert proc.returncode == 0, proc.stderr
    summary = json.loads(proc.stdout)
    assert summary["ok"] is True
    assert summary["nonempty_paragraphs"] == 2
    text = out.read_text(encoding="utf-8")
    assert "制定本办法" in text
    assert "[P0]" in text
    assert "[P2]" in text
